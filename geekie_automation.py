from __future__ import annotations

import argparse
import csv
import os
import re
import shutil
import sys
import time
import unicodedata
from pathlib import Path

from openpyxl import load_workbook
from selenium import webdriver
from selenium.common.exceptions import StaleElementReferenceException, TimeoutException
from selenium.webdriver import ChromeOptions
from selenium.webdriver.common.action_chains import ActionChains
from selenium.webdriver.common.by import By
from selenium.webdriver.common.keys import Keys
from selenium.webdriver.support import expected_conditions as EC
from selenium.webdriver.support.ui import WebDriverWait


LAB_URL = "https://educadores.geekieteste.geekie.com.br/"
ONE_URL = "https://one.geekie.com.br/"
PASSWORD = os.environ.get("GEEKIE_PASSWORD", "REMOVIDO")
PROFILE = Path(__file__).with_name("chrome_profile")
CHROME_USER_DATA = Path.home() / "AppData/Local/Google/Chrome/User Data"
ERRORS = Path(__file__).with_name("erros.csv")


class AccountNotFound(RuntimeError):
    pass


class EmailAlreadyExists(RuntimeError):
    pass


def wait(driver, seconds=15):
    return WebDriverWait(driver, seconds)


def normalize_text(value):
    text = "" if value is None else str(value)
    # Remove acentos
    text = unicodedata.normalize("NFKD", text)
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    # Remove caracteres especiais (apÃ³strofos, traÃ§os, etc) e espaÃ§os extras
    text = re.sub(r"[^a-zA-Z0-9\s]", "", text)
    return re.sub(r"\s+", " ", text).strip().lower()


def click_text(driver, text, seconds=15):
    xpath = f"//*[self::button or @role='button' or self::a][normalize-space()={text!r}]"
    element = wait(driver, seconds).until(EC.element_to_be_clickable((By.XPATH, xpath)))
    driver.execute_script("arguments[0].click()", element)
    return element


def new_driver():
    options = ChromeOptions()
    options.add_argument(f"--user-data-dir={PROFILE.resolve()}")
    options.add_argument("--profile-directory=Default")
    options.add_argument("--start-maximized")
    # Bloquear popups de notificaÃ§Ã£o do navegador automaticamente
    options.add_argument("--disable-notifications")
    prefs = {
        "profile.default_content_setting_values.notifications": 2,
        "profile.default_content_setting_values.automatic_downloads": 1,
    }
    options.add_experimental_option("prefs", prefs)
    options.add_experimental_option("excludeSwitches", ["enable-automation"])
    return webdriver.Chrome(options=options)


def clone_chrome_profile():
    source = CHROME_USER_DATA
    if not source.exists():
        raise RuntimeError(f"Perfil do Chrome nÃ£o encontrado: {source}")
    if PROFILE.exists():
        shutil.rmtree(PROFILE)
    PROFILE.mkdir(parents=True)
    shutil.copy2(source / "Local State", PROFILE / "Local State")
    shutil.copytree(
        source / "Default",
        PROFILE / "Default",
        ignore=shutil.ignore_patterns("Cache", "Code Cache", "GPUCache", "Service Worker"),
    )
    print("SessÃ£o do Chrome copiada.")


def setup_profile():
    driver = new_driver()
    driver.get(LAB_URL)
    driver.execute_script("window.open(arguments[0], '_blank')", ONE_URL)
    input("Entre no Geekie Lab e Geekie One nas abas abertas. Depois pressione ENTER aqui...")
    driver.quit()


def ensure_logins(driver):
    driver.get(LAB_URL)
    if driver.find_elements(By.NAME, "username_or_email"):
        input("FaÃ§a login no Geekie Lab nesta janela e pressione ENTER aqui...")
        driver.get(LAB_URL)
    if driver.find_elements(By.NAME, "username_or_email"):
        raise RuntimeError("Geekie Lab continua deslogado")

    driver.get(ONE_URL)
    time.sleep(1)
    if driver.find_elements(By.XPATH, "//*[normalize-space()='Com uma conta Geekie One']"):
        input("Se necessÃ¡rio, faÃ§a login no Geekie One e pressione ENTER aqui...")


def log_error(row, name, reason):
    exists = ERRORS.exists()
    with ERRORS.open("a", newline="", encoding="utf-8") as file:
        writer = csv.writer(file)
        if not exists:
            writer.writerow(["linha", "aluno", "erro"])
        writer.writerow([row, name, reason])


def normalize_ascii(value):
    text = unicodedata.normalize("NFKD", str(value))
    return "".join(ch for ch in text if not unicodedata.combining(ch))


def username_from_email(email, used_usernames):
    local = str(email).split("@", 1)[0].lower()
    local = re.sub(r"-escola.*$", "", local)
    base = re.sub(r"[^a-z0-9]", "", normalize_ascii(local))
    if not base:
        raise RuntimeError(f"email invalido para gerar usuario: {email}")
    used = {normalize_text(item) for item in used_usernames if item}
    if normalize_text(base) not in used:
        return base
    for number in range(2, 100):
        candidate = f"{base}{number}"
        if normalize_text(candidate) not in used:
            return candidate
    raise RuntimeError(f"nao consegui gerar usuario livre para {email}")


def dismiss_notification_popup(driver):
    """Fecha o popup de notificaÃ§Ã£o do navegador, se presente."""
    try:
        btn = driver.find_element(By.XPATH, "//*[normalize-space()='NÃ£o, obrigado' or normalize-space()='Nao, obrigado']")
        if btn.is_displayed():
            driver.execute_script("arguments[0].click()", btn)
            time.sleep(0.5)
    except Exception:
        pass
    # TambÃ©m tenta fechar via JavaScript caso seja um prompt nativo
    try:
        driver.execute_script(
            "if (window.Notification) { Notification.requestPermission = function() { return Promise.resolve('denied'); }; }"
        )
    except Exception:
        pass


def fill_field(driver_or_field, value_or_field=None, value=None):
    if value is None:
        field = driver_or_field
        val = str(value_or_field)
        driver = None
    else:
        driver = driver_or_field
        field = value_or_field
        val = str(value)

    # 1. Tenta via teclado (padrÃ£o)
    try:
        field.click()
        field.send_keys(Keys.CONTROL, "a")
        field.send_keys(Keys.BACKSPACE)
        field.send_keys(val)
        time.sleep(0.2)
    except Exception:
        pass

    # 2. Se falhar ou se tivermos driver, usa a tÃ©cnica de injeÃ§Ã£o JavaScript (React/Angular)
    if driver is not None:
        try:
            driver.execute_script("""
                var el = arguments[0];
                var val = arguments[1];
                
                // ForÃ§a foco e clique
                el.focus();
                el.click();
                
                // InjeÃ§Ã£o via protÃ³tipo (engana o tracker do React)
                var setter = Object.getOwnPropertyDescriptor(window.HTMLInputElement.prototype, 'value').set;
                setter.call(el, val);
                
                // Dispara eventos nativos para o framework detectar a mudanÃ§a
                el.dispatchEvent(new InputEvent('input', { bubbles: true, inputType: 'insertText', data: val }));
                el.dispatchEvent(new Event('input', { bubbles: true }));
                el.dispatchEvent(new Event('change', { bubbles: true }));
                el.dispatchEvent(new Event('blur', { bubbles: true }));
                
                // Fallback direto
                if (el.value !== val) { el.value = val; }
            """, field, val)
            time.sleep(0.3)
        except Exception:
            pass


def class_parts(class_code):
    text = normalize_ascii(str(class_code)).upper().strip()
    letter_match = re.search(r"([A-Z])$", text)
    if not letter_match:
        raise RuntimeError(f"turma invalida: {class_code}")
    letter = letter_match.group(1)
    if re.search(r"(EM)?3[A-Z]?$", text) or "3" in text:
        return "3", letter, f"3Âª SÃ©rie {letter}"
    if "9" in text:
        return "9", letter, f"9Âº {letter}"
    number_match = re.search(r"(\d+)", text)
    if number_match:
        grade = number_match.group(1)
        return grade, letter, f"{grade} {letter}"
    raise RuntimeError(f"turma invalida: {class_code}")


def ensure_one_logged_out(driver):
    driver.get(ONE_URL)
    time.sleep(2)
    # A versao nova abre diretamente nos campos de usuario e senha.
    try:
        one_login_fields(driver, 2)
        return
    except TimeoutException:
        pass
    if driver.find_elements(By.XPATH, "//*[normalize-space()='Com uma conta Geekie One']"):
        return
    try:
        profile = wait(driver, 4).until(EC.element_to_be_clickable((By.CSS_SELECTOR, "[aria-label='Perfil']")))
        driver.execute_script("arguments[0].click()", profile)
        click_text(driver, "Sair", 8)
        wait(driver, 20).until(lambda d: (
            bool(d.find_elements(By.XPATH, "//*[normalize-space()='Com uma conta Geekie One']"))
            or bool([field for field in d.find_elements(By.CSS_SELECTOR, "input") if field.is_displayed() and normalize_text(field.get_attribute("type")) == "password"])
        ))
        return
    except TimeoutException:
        driver.execute_script(
            """
            localStorage.clear();
            sessionStorage.clear();
            document.cookie.split(';').forEach(function(cookie) {
              document.cookie = cookie.replace(/^ +/, '').replace(/=.*/, '=;expires=' + new Date(0).toUTCString() + ';path=/');
            });
            // Limpeza profunda de IndexedDB
            if (window.indexedDB && window.indexedDB.databases) {
                window.indexedDB.databases().then(dbs => {
                    dbs.forEach(db => window.indexedDB.deleteDatabase(db.name));
                });
            }
            """
        )
        driver.delete_all_cookies()
        driver.get(ONE_URL)
        try:
            wait(driver, 30).until(lambda d: (
                bool(d.find_elements(By.XPATH, "//*[normalize-space()='Com uma conta Geekie One']"))
                or bool([field for field in d.find_elements(By.CSS_SELECTOR, "input") if field.is_displayed() and normalize_text(field.get_attribute("type")) == "password"])
            ))
        except TimeoutException:
            raise RuntimeError("nao consegui deslogar do Geekie One antes de trocar de aluno")


def reset_password(driver, link, name):
    driver.get(link)
    password_fields = wait(driver, 20).until(
        lambda d: [
            field for field in d.find_elements(
                By.CSS_SELECTOR, "input[type='password'], input[name*='password']"
            )
            if field.is_displayed()
        ]
    )
    if len(password_fields) < 2:
        raise RuntimeError("campos de senha do reset nao encontrados")
    fill_field(driver, password_fields[0], PASSWORD)
    time.sleep(0.3)
    fill_field(driver, password_fields[1], PASSWORD)
    time.sleep(0.3)
    # Verificar se os campos foram preenchidos
    v1 = password_fields[0].get_attribute("value") or ""
    v2 = password_fields[1].get_attribute("value") or ""
    if not v1 or not v2:
        raise RuntimeError(f"senha nao preenchida nos campos de reset (v1={len(v1)} chars, v2={len(v2)} chars)")
    
    # AutomatizaÃ§Ã£o do clique no botÃ£o de Redefinir/Salvar
    print(f"Tentando redefinir senha automaticamente para {name}...")
    
    clicked = False
    # 1. Tenta por seletores especÃ­ficos do HTML inspecionado pelo usuÃ¡rio
    selectors = [
        "input[type='submit'][value='REDEFINIR']",
        "input.geekieui-btn-temporary",
        "button[type='submit']",
        "input[type='submit']"
    ]
    
    for selector in selectors:
        try:
            btn = driver.find_element(By.CSS_SELECTOR, selector)
            if btn.is_displayed():
                driver.execute_script("arguments[0].click();", btn)
                print(f"BotÃ£o clicado via seletor: {selector}")
                clicked = True
                break
        except Exception:
            continue
            
    # 2. Se falhar, tenta pelos textos conhecidos
    if not clicked:
        reset_buttons = ["Redefinir senha", "Salvar", "Confirmar", "Redefinir", "REDEFINIR"]
        for btn_text in reset_buttons:
            try:
                click_text(driver, btn_text, 2)
                clicked = True
                print(f"BotÃ£o '{btn_text}' clicado.")
                break
            except TimeoutException:
                continue
            
    if not clicked:
        print(f"Aviso: nÃ£o encontrei o botÃ£o de redefinir automaticamente. Por favor, clique manualmente no Chrome.")
        input("Pressione ENTER aqui apÃ³s clicar no botÃ£o no Chrome...")
    else:
        # Espera um pouco para a redefiniÃ§Ã£o ser processada
        time.sleep(3)
        print("Senha redefinida com sucesso.")


def one_login_fields(driver, seconds=20):
    def fields(d):
        visible = [field for field in d.find_elements(By.CSS_SELECTOR, "input") if field.is_displayed()]
        password = [field for field in visible if normalize_text(field.get_attribute("data-testid")) == "password-input"]
        if not password:
            password = [field for field in visible if normalize_text(field.get_attribute("type")) == "password"]
        user = [
            field for field in visible
            if normalize_text(field.get_attribute("data-testid")) in ("user-input", "username-input", "email-input")
        ]
        if not user:
            user = [
                field for field in visible
            if any(token in normalize_text(" ".join(
                field.get_attribute(name) or "" for name in ("placeholder", "aria-label", "name")
            )) for token in ("usuario", "email"))
            ]
        if user and password:
            return [user[0], password[0]]
        text_inputs = [field for field in visible if normalize_text(field.get_attribute("type")) in ("", "text")]
        if text_inputs and password:
            return [text_inputs[0], password[0]]
        return False
    return wait(driver, seconds).until(fields)


def open_one_login(driver):
    """Aceita a tela antiga com etapa intermediaria e a tela nova ja direta."""
    dismiss_notification_popup(driver)
    try:
        click_text(driver, "Usar Aqui", 2)
    except TimeoutException:
        pass
    try:
        return one_login_fields(driver, 2)
    except TimeoutException:
        click_text(driver, "Com uma conta Geekie One", 10)
        return one_login_fields(driver, 20)


def login_one(driver, login_id, used_usernames=None):
    ensure_one_logged_out(driver)
    driver.get(ONE_URL)
    visible = open_one_login(driver)
    if len(visible) < 2:
        raise RuntimeError("campos de login do Geekie One nao encontrados")
    # Usa o login_id (que agora serÃ¡ o e-mail)
    user_field, pw_field = visible[0], visible[1]

    def fill_and_verify(field, value, label):
        for _ in range(3):
            fill_field(driver, field, value)
            current = field.get_attribute("value") or ""
            if current == str(value):
                return
            try:
                ActionChains(driver).move_to_element(field).click().key_down(Keys.CONTROL).send_keys("a").key_up(Keys.CONTROL).send_keys(str(value)).perform()
                if (field.get_attribute("value") or "") == str(value):
                    return
            except Exception:
                pass
            try:
                if label == "usuario":
                    field = wait(driver, 3).until(EC.visibility_of_element_located((
                        By.CSS_SELECTOR,
                        "input[data-testid='user-input'], input[data-testid='username-input'], input[data-testid='email-input'], input[placeholder*='usuário'], input[placeholder*='usuario']",
                    )))
                else:
                    field = wait(driver, 3).until(EC.visibility_of_element_located((By.CSS_SELECTOR, "input[data-testid='password-input'], input[type='password']")))
            except TimeoutException:
                pass
        raise RuntimeError(f"campo de {label} nao foi preenchido no Geekie One")

    fill_and_verify(user_field, login_id, "usuario")
    fill_and_verify(pw_field, PASSWORD, "senha")
    # Texto exato: evita confundir o login normal com Google, Microsoft ou Apple.
    click_text(driver, "Acessar", 10)
    start = time.time()
    invalid_markers = (
        "usuariosenha invalidos",
        "usuario senha invalidos",
        "usuario ou senha invalidos",
        "senha invalida",
        "senha invalid",
        "credenciais invalid",
    )
    while time.time() - start < 45:
        text = normalize_text(normalize_ascii(driver.execute_script("return document.body?.innerText || ''")))
        if any(marker in text for marker in invalid_markers) or ("invalid" in text and ("senha" in text or "usuario" in text)):
            raise RuntimeError("Usuario/Senha invalidos no Geekie One")
        visible_inputs = [field for field in driver.find_elements(By.CSS_SELECTOR, "input") if field.is_displayed()]
        login_button = driver.find_elements(By.XPATH, "//*[normalize-space()='Acessar']")
        still_login = (
            "acessar com uma conta geekie one" in text
            or ("nome de usuario ou email" in text and "senha" in text)
            or (len(visible_inputs) >= 2 and bool(login_button))
        )
        if not still_login and "aproveite o geekie one" not in text:
            break
        if time.time() - start > 10 and still_login:
            raise RuntimeError("Geekie One nao logou; possivel Usuario/Senha invalidos")
        time.sleep(0.5)
    else:
        raise RuntimeError("timeout ao entrar no Geekie One")
    return complete_profile_after_login(driver, login_id, used_usernames)


def click_enabled_contains(driver, text, seconds=10):
    target = normalize_text(text)
    deadline = time.time() + seconds
    while time.time() < deadline:
        for element in reversed(driver.find_elements(By.XPATH, "//*[self::button or @role='button' or self::a]")):
            try:
                disabled = element.get_attribute("disabled")
                aria_disabled = element.get_attribute("aria-disabled")
                if element.is_displayed() and not disabled and aria_disabled != "true" and target in normalize_text(element.text):
                    driver.execute_script("arguments[0].click()", element)
                    return True
            except StaleElementReferenceException:
                continue
        time.sleep(0.25)
    return False


def cdp_click(driver, x, y):
    driver.execute_cdp_cmd(
        "Input.dispatchMouseEvent",
        {"type": "mousePressed", "x": x, "y": y, "button": "left", "clickCount": 1},
    )
    driver.execute_cdp_cmd(
        "Input.dispatchMouseEvent",
        {"type": "mouseReleased", "x": x, "y": y, "button": "left", "clickCount": 1},
    )
    time.sleep(0.2)


def cdp_type_text(driver, text):
    driver.execute_cdp_cmd("Input.insertText", {"text": str(text)})
    time.sleep(0.4)


def cdp_eval(driver, expression):
    return driver.execute_cdp_cmd(
        "Runtime.evaluate",
        {
            "expression": expression,
            "awaitPromise": True,
            "returnByValue": True,
        },
    ).get("result", {}).get("value")


def fill_profile_update_by_cdp_dom(driver, username):
    script = f"""
    (() => {{
      const username = {username!r};
      const norm = (value) => (value || '')
        .normalize('NFD').replace(/[\\u0300-\\u036f]/g, '')
        .replace(/[^a-zA-Z0-9\\s]/g, '')
        .replace(/\\s+/g, ' ').trim().toLowerCase();
      const all = [...document.querySelectorAll('input, textarea')];
      const userInput = all.find((el) => norm(el.getAttribute('placeholder')).includes('usu'));
      if (!userInput) return 'NO_USER_INPUT:' + all.length;
      userInput.focus();
      userInput.click();
      const setter = Object.getOwnPropertyDescriptor(HTMLInputElement.prototype, 'value')?.set
        || Object.getOwnPropertyDescriptor(HTMLTextAreaElement.prototype, 'value')?.set;
      setter.call(userInput, username);
      userInput.dispatchEvent(new InputEvent('input', {{bubbles: true, inputType: 'insertText', data: username}}));
      userInput.dispatchEvent(new Event('change', {{bubbles: true}}));
      userInput.dispatchEvent(new Event('blur', {{bubbles: true}}));

      const checkButton = [...document.querySelectorAll('button, [role=button], a')]
        .find((el) => norm(el.innerText || el.textContent).includes('checar'));
      if (checkButton) checkButton.click();

      const noCpf = [...document.querySelectorAll('label, span, div, p')]
        .find((el) => norm(el.innerText || el.textContent).includes('nao quero informar o cpf'));
      if (noCpf) noCpf.click();
      else {{
        const radios = [...document.querySelectorAll('input[type=radio]')];
        if (radios.length) radios[radios.length - 1].click();
      }}
      return 'OK';
    }})()
    """
    deadline = time.time() + 15
    last_result = None
    while time.time() < deadline:
        last_result = cdp_eval(driver, script)
        if last_result == "OK":
            time.sleep(1.2)
            return
        time.sleep(0.5)
    raise RuntimeError(f"CDP nao encontrou/preencheu usuario: {last_result}")


def wait_profile_update_inputs(driver, seconds=15):
    deadline = time.time() + seconds
    while time.time() < deadline:
        result = cdp_eval(
            driver,
            """
            (() => {
              const norm = (value) => (value || '')
                .normalize('NFD').replace(/[\\u0300-\\u036f]/g, '')
                .replace(/[^a-zA-Z0-9\\s]/g, '')
                .replace(/\\s+/g, ' ').trim().toLowerCase();
              const inputs = [...document.querySelectorAll('input, textarea')];
              const username = inputs.find((el) => norm(el.getAttribute('placeholder')).includes('usu'));
              return username ? 'READY' : 'WAIT:' + inputs.length;
            })()
            """,
        )
        if result == "READY":
            return True
        time.sleep(0.5)
    return False


def click_save_profile_by_cdp_dom(driver):
    deadline = time.time() + 20
    while time.time() < deadline:
        result = cdp_eval(
            driver,
            """
            (() => {
              const norm = (value) => (value || '')
                .normalize('NFD').replace(/[\\u0300-\\u036f]/g, '')
                .replace(/[^a-zA-Z0-9\\s]/g, '')
                .replace(/\\s+/g, ' ').trim().toLowerCase();
              const buttons = [...document.querySelectorAll('button, [role=button], a, div, span')]
                .filter((el) => norm(el.innerText || el.textContent).includes('salvar'));
              for (const btn of buttons.reverse()) {
                if (!btn.disabled && btn.getAttribute('aria-disabled') !== 'true') {
                  btn.click();
                  return 'OK';
                }
              }
              return 'WAIT';
            })()
            """,
        )
        if result == "OK":
            time.sleep(2)
            return True
        time.sleep(0.5)
    return False


def fill_profile_update_by_coordinates(driver, username):
    size = driver.execute_script("return {w: window.innerWidth, h: window.innerHeight};")
    field_x = int(size["w"] * 0.50)
    check_x = int(size["w"] * 0.60)
    radio_x = int(size["w"] * 0.375)
    username_y = int(size["h"] * 0.655)
    cpf_y = int(size["h"] * 0.807)
    save_y = int(size["h"] * 0.885)
    cdp_click(driver, field_x, username_y)
    cdp_type_text(driver, username)
    time.sleep(0.4)
    cdp_click(driver, check_x, username_y)
    time.sleep(1)
    cdp_click(driver, radio_x, cpf_y)
    time.sleep(0.4)
    cdp_click(driver, field_x, save_y)
    time.sleep(2)


def select_no_cpf(driver):
    def cpf_radio_point():
        return driver.execute_script(
            """
            const norm = (value) => (value || '')
              .normalize('NFD').replace(/[\\u0300-\\u036f]/g, '')
              .replace(/[^a-zA-Z0-9\\s]/g, '')
              .replace(/\\s+/g, ' ').trim().toLowerCase();
            const visible = (el) => {
              const r = el.getBoundingClientRect();
              const st = getComputedStyle(el);
              return r.width > 0 && r.height > 0 && st.visibility !== 'hidden' && st.display !== 'none';
            };
            const textSpan = [...document.querySelectorAll('span')]
              .find((el) => visible(el) && norm(el.innerText || el.textContent) === 'nao quero informar o cpf doa estudante neste momento');
            const radio = textSpan?.closest('[role="radio"]');
            if (!radio) return null;
            radio.scrollIntoView({block:'center'});
            const box = radio.querySelector('[data-testid="radiobox-undefined"]') || radio.firstElementChild || radio;
            const r = box.getBoundingClientRect();
            return {x:r.left + r.width / 2, y:r.top + r.height / 2};
            """
        )

    for _ in range(5):
        point = cpf_radio_point()
        if point:
            cdp_click(driver, int(point["x"]), int(point["y"]))
            time.sleep(0.7)
            checked = driver.execute_script(
                "return !!document.querySelector('[role=\"radio\"][aria-checked=\"true\"], input[type=radio]:checked');"
            )
            if checked:
                print("CPF marcado via coordenada da bolinha.")
                return True
    for _ in range(3):
        try:
            radio = wait(driver, 2).until(EC.presence_of_element_located((
                By.XPATH,
                "//*[@role='radio' and contains(normalize-space(.), 'CPF') and contains(normalize-space(.), 'informar')]",
            )))
            driver.execute_script("arguments[0].scrollIntoView({block:'center'});", radio)
            try:
                box = radio.find_element(By.CSS_SELECTOR, "[data-testid='radiobox-undefined']")
            except Exception:
                box = radio
            ActionChains(driver).move_to_element(box).click().perform()
            time.sleep(0.7)
            checked = driver.execute_script(
                "return arguments[0].getAttribute('aria-checked') === 'true' || !!arguments[0].querySelector('[aria-checked=\"true\"], input[type=radio]:checked');",
                radio,
            )
            if checked:
                print("CPF marcado via ActionChains na bolinha.")
                return True
            ActionChains(driver).move_to_element(radio).click().perform()
            time.sleep(0.7)
            checked = driver.execute_script(
                "return arguments[0].getAttribute('aria-checked') === 'true' || !!arguments[0].querySelector('[aria-checked=\"true\"], input[type=radio]:checked');",
                radio,
            )
            if checked:
                print("CPF marcado via ActionChains no radio.")
                return True
        except Exception:
            pass
    for _ in range(4):
        result = driver.execute_script(
            """
            const norm = (value) => (value || '')
              .normalize('NFD').replace(/[\\u0300-\\u036f]/g, '')
              .replace(/[^a-zA-Z0-9\\s]/g, '')
              .replace(/\\s+/g, ' ').trim().toLowerCase();
            const visible = (el) => {
              const r = el.getBoundingClientRect();
              const st = getComputedStyle(el);
              return r.width > 0 && r.height > 0 && st.visibility !== 'hidden' && st.display !== 'none';
            };
            const radios = [...document.querySelectorAll('[role="radio"], input[type=radio]')].filter(visible);
            let radio = radios.find((el) => norm([el.getAttribute('aria-label'), el.innerText, el.textContent].join(' ')).includes('nao quero informar o cpf'));
            if (!radio) {
              const label = [...document.querySelectorAll('label, span, div, p')]
                .find((el) => visible(el) && norm(el.innerText || el.textContent).includes('nao quero informar o cpf'));
              radio = label?.closest('[role="radio"]') || label?.querySelector('[role="radio"], input[type=radio]') || label;
            }
            if (!radio && radios.length) radio = radios[radios.length - 1];
            if (!radio) return {ok:false, checked:false};
            const target = radio.closest('[role="radio"]') || radio.closest('label') || radio;
            target.scrollIntoView({block:'center'});
            const explicitBox = target.querySelector('[data-testid="radiobox-undefined"]') || document.querySelector('[data-testid="radiobox-undefined"]');
            const clickTarget = explicitBox || [...target.querySelectorAll('*')]
              .filter(visible)
              .map((el) => ({el, r: el.getBoundingClientRect()}))
              .filter(({r}) => r.width >= 8 && r.width <= 40 && r.height >= 8 && r.height <= 40)
              .sort((a, b) => (a.r.width * a.r.height) - (b.r.width * b.r.height))[0]?.el || target;
            for (const type of ['pointerdown', 'mousedown', 'mouseup', 'click']) {
              clickTarget.dispatchEvent(new MouseEvent(type, {bubbles:true, cancelable:true, view:window}));
            }
            clickTarget.click();
            const checked = target.getAttribute('aria-checked') === 'true'
              || radio.getAttribute('aria-checked') === 'true'
              || radio.checked === true
              || !!target.querySelector('[aria-checked="true"], input[type=radio]:checked');
            const r = clickTarget.getBoundingClientRect();
            return {ok:true, checked, x:r.left + r.width / 2, y:r.top + r.height / 2};
            """
        )
        if result and result.get("checked"):
            time.sleep(0.5)
            print("CPF marcado via JS.")
            return True
        if result and result.get("ok"):
            cdp_click(driver, int(result["x"]), int(result["y"]))
            time.sleep(0.7)
            checked = driver.execute_script(
                """
                return !!document.querySelector('[role="radio"][aria-checked="true"], input[type=radio]:checked');
                """
            )
            if checked:
                print("CPF marcado via CDP.")
                return True
    coords = driver.execute_script(
        """
        const norm = (value) => (value || '')
          .normalize('NFD').replace(/[\\u0300-\\u036f]/g, '')
          .replace(/[^a-zA-Z0-9\\s]/g, '')
          .replace(/\\s+/g, ' ').trim().toLowerCase();
        const visible = (el) => {
          const r = el.getBoundingClientRect();
          const st = getComputedStyle(el);
          return r.width > 0 && r.height > 0 && st.visibility !== 'hidden' && st.display !== 'none';
        };
        const label = [...document.querySelectorAll('span, div, p, label')]
          .find((el) => visible(el) && norm(el.innerText || el.textContent).includes('nao quero informar o cpf'));
        if (!label) return null;
        label.scrollIntoView({block:'center'});
        const r = label.getBoundingClientRect();
        return {x: Math.max(8, r.left - 18), y: r.top + r.height / 2};
        """
    )
    if coords:
        for _ in range(3):
            cdp_click(driver, int(coords["x"]), int(coords["y"]))
            time.sleep(0.6)
            checked = driver.execute_script(
                "return !!document.querySelector('[role=\"radio\"][aria-checked=\"true\"], input[type=radio]:checked');"
            )
            if checked:
                print("CPF marcado via fallback por texto.")
                return True
    return False


def click_save_profile(driver):
    deadline = time.time() + 20
    while time.time() < deadline:
        clicked = driver.execute_script(
            """
            const norm = (value) => (value || '')
              .normalize('NFD').replace(/[\\u0300-\\u036f]/g, '')
              .replace(/[^a-zA-Z0-9\\s]/g, '')
              .replace(/\\s+/g, ' ').trim().toLowerCase();
            const visible = (el) => {
              const r = el.getBoundingClientRect();
              const st = getComputedStyle(el);
              return r.width > 0 && r.height > 0 && st.visibility !== 'hidden' && st.display !== 'none';
            };
            const buttons = [...document.querySelectorAll('button, [role=button], a, div, span')]
              .filter((el) => visible(el) && norm(el.innerText || el.textContent).includes('salvar'));
            for (const btn of buttons.reverse()) {
              const disabled = btn.disabled || btn.getAttribute('aria-disabled') === 'true';
              if (!disabled) {
                btn.click();
                return true;
              }
            }
            return false;
            """
        )
        if clicked:
            return True
        time.sleep(0.5)
    return False


def complete_profile_after_login(driver, email, used_usernames=None):
    visible_page = normalize_text(normalize_ascii(driver.execute_script("return document.body?.innerText || ''")))
    has_profile_update = (
        "atualizacao das informacoes" in visible_page
        and ("nome de usuario" in visible_page or "cpf" in visible_page or "salvar alteracoes" in visible_page)
    )
    if not has_profile_update:
        return None

    username = None
    base_username = username_from_email(email, used_usernames or [])
    inputs_ready = wait_profile_update_inputs(driver, 15)
    if not inputs_ready:
        if "cpf" in visible_page and "salvar alteracoes" in visible_page and "nome de usuario" not in visible_page:
            print("Atualizacao sem campo de usuario; preenchendo apenas CPF/status.")
            select_no_cpf(driver)
            if not (click_save_profile_by_cdp_dom(driver) or click_save_profile(driver)):
                raise RuntimeError("nao consegui salvar a atualizacao das informacoes")
            time.sleep(2)
            return None
    try:
        if inputs_ready:
            username = ensure_available_signup_username(driver, base_username)
            print(f"Usuario preenchido na atualizacao: {username}")
        else:
            raise RuntimeError("inputs da atualizacao ainda nao apareceram")
    except Exception as exc:
        print(f"Aviso: tentativa normal falhou; usando campo de usuario por posicao: {exc}")
        username = base_username
        try:
            field = visible_input_by_index(driver, 1)
            fill_field(driver, field, username)
            click_if_present(driver, "CHECAR", 2)
            time.sleep(0.8)
        except Exception as position_exc:
            print(f"Aviso: fallback por input falhou; usando CDP DOM: {position_exc}")
            try:
                fill_profile_update_by_cdp_dom(driver, username)
                print(f"Usuario preenchido via CDP DOM: {username}")
            except Exception as cdp_exc:
                print(f"Aviso: fallback CDP DOM falhou; usando coordenadas: {cdp_exc}")
                fill_profile_update_by_coordinates(driver, username)

    select_no_cpf(driver)

    if not (click_save_profile_by_cdp_dom(driver) or click_save_profile(driver)):
        raise RuntimeError("nao consegui salvar a atualizacao das informacoes")
    time.sleep(2)
    return username


def complete_profile_if_needed(driver):
    try:
        radio = wait(driver, 4).until(EC.element_to_be_clickable((By.XPATH, "//*[contains(., 'NÃ£o quero informar o CPF') and @role='radio']")))
        driver.execute_script("arguments[0].click()", radio)
        click_text(driver, "Salvar alteraÃ§Ãµes")
    except TimeoutException:
        pass


def activate(driver, code, class_code):
    complete_profile_if_needed(driver)
    # Tenta fechar eventuais popups iniciais que bloqueiam a tela
    for text in ("Pular", "Agora nÃ£o", "NÃ£o agora", "Entendi", "Fechar"):
        click_if_present(driver, text, 2)
    
    # Simplificado: espera o botÃ£o de ativaÃ§Ã£o ou o container principal, sem travar em textos ocultos
    print("Aguardando tela do Geekie One estar pronta...")
    visible_text = normalize_text(normalize_ascii(driver.execute_script("return document.body?.innerText || ''")))
    if "acessar com uma conta geekie one" in visible_text or ("nome de usuario ou email" in visible_text and "senha" in visible_text):
        raise RuntimeError("nao esta logado no Geekie One; tela de login aberta")
    if "atualizacao das informacoes" in visible_text:
        raise RuntimeError("tela de atualizacao das informacoes aberta; complete o perfil antes de ativar")
    if "material padrao 2026" in visible_text or "material de 2026 ja esta aqui" in visible_text:
        raise RuntimeError("material 2026 ja esta aberto; nao deve ativar novo codigo")
    try:
        # Espera curta para ver se o botÃ£o jÃ¡ estÃ¡ lÃ¡ (caso do screenshot do usuÃ¡rio)
        wait(driver, 25).until(EC.element_to_be_clickable((By.XPATH, "//*[self::button or @role='button' or self::a][.//*[normalize-space()='Ativar novo material'] or normalize-space()='Ativar novo material']")))
    except TimeoutException:
        # Se nÃ£o estiver, espera o sumiÃ§o da frase de carregamento, mas com timeout menor e fallback
        try:
            wait(driver, 30).until(lambda d: "Aproveite o Geekie One" not in d.page_source)
        except TimeoutException:
            pass
        
    # Tenta clicar no botÃ£o de ativaÃ§Ã£o com vÃ¡rias tentativas
    activated = False
    for _ in range(3):
        try:
            click_text(driver, "Ativar novo material", 20)
            activated = True
            break
        except TimeoutException:
            # Tenta fechar popups novamente se falhar
            click_if_present(driver, "Pular", 2)
            driver.refresh()
            time.sleep(2)
    
    if not activated:
        raise RuntimeError("nÃ£o consegui encontrar ou clicar no botÃ£o 'Ativar novo material'")

    field = wait(driver).until(EC.visibility_of_element_located((By.CSS_SELECTOR, "input[type='text']")))
    field.send_keys(code)
    click_text(driver, "Ativar")
    click_text(driver, "Continuar", 10)
    click_text(driver, "Confirmar", 20)
    print("Material ativado. Aguardando tela de acesso ao material...")
    
    # Tenta clicar no botÃ£o de acesso, mas sem travar em "Carregando" oculto
    try:
        click_text(driver, "Acessar seu material Geekie One!", 60)
    except TimeoutException:
        # Se falhar, tenta um refresh e busca novamente
        driver.refresh()
        click_text(driver, "Acessar seu material Geekie One!", 30)
    
    print("Aguardando carregamento da lista de turmas...")
    select_xpath = "//input[@readonly and (contains(@value, 'turma') or contains(@class, 'input'))]"
    select = wait(driver, 60).until(EC.element_to_be_clickable((By.XPATH, select_xpath)))
    current_val = select.get_attribute("value") or ""
    grade_token, letter, target_class = class_parts(class_code)

    if letter not in current_val:
        print(f"Turma atual '{current_val}' diferente de '{letter}'. ForÃ§ando troca.")
        
        # Tenta abrir o menu de vÃ¡rias formas
        opened = False
        # Adicionado 'Overlay click' como prioridade, pois o erro mostrou que ele bloqueia o clique
        for method in ["Overlay click", "JS click", "Native click", "Space key", "Container click"]:
            try:
                if method == "Overlay click":
                    # Busca o overlay que o erro do usuÃ¡rio mostrou estar na frente
                    overlays = driver.find_elements(By.CSS_SELECTOR, "[data-testid='DropdownInput-overlay']")
                    if overlays:
                        driver.execute_script("arguments[0].click()", overlays[0])
                    else:
                        continue
                elif method == "JS click":
                    driver.execute_script("arguments[0].click()", select)
                elif method == "Native click":
                    select.click()
                elif method == "Space key":
                    select.send_keys(Keys.SPACE)
                elif method == "Container click":
                    parent = select.find_element(By.XPATH, "./..")
                    driver.execute_script("arguments[0].click()", parent)
                
                time.sleep(2)
                # Verifica se alguma opÃ§Ã£o de turma apareceu
                if driver.find_elements(By.XPATH, f"//*[(@role='option' or @role='button' or self::li) and contains(., {grade_token!r})]"):
                    print(f"Menu aberto com sucesso via: {method}")
                    opened = True
                    break
            except Exception as e:
                print(f"Falha ao tentar abrir menu via {method}: {e}")
                continue
        
        # Screenshot de depuraÃ§Ã£o para ver se o menu abriu
        driver.save_screenshot(f"debug_menu_aberto_linha_{getattr(driver, 'current_row_num', 'unknown')}.png")
        
        if not opened:
            print("Aviso: nÃ£o detectei a abertura do menu visualmente, mas tentarei clicar na opÃ§Ã£o assim mesmo.")
            
        print(f"Selecionando turma: {target_class}")
    else:
        print(f"Turma '{letter}' jÃ¡ parece estar selecionada.")
        click_text(driver, "Selecionar turma")
        return
    
    # XPath flexÃ­vel para encontrar a opÃ§Ã£o na lista
    # Procura por elementos que tenham '9' e a letra, priorizando os que parecem ser itens de lista/botÃµes
    option_selectors = [
        f"//*[(@role='option' or @role='button') and contains(., {grade_token!r}) and contains(., {letter!r})]",
        f"//li[contains(., {grade_token!r}) and contains(., {letter!r})]",
        f"//div[contains(., {grade_token!r}) and contains(., {letter!r}) and not(contains(@class, 'input'))]"
    ]
    
    option = None
    for selector in option_selectors:
        try:
            elements = driver.find_elements(By.XPATH, selector)
            visible = [el for el in elements if el.is_displayed() and el.size['width'] > 0]
            if visible:
                option = visible[0]
                break
        except Exception:
            continue
            
    if not option:
        # Ãšltimo recurso: busca qualquer texto visÃ­vel que contenha a turma
        try:
            option = wait(driver, 10).until(lambda d: next((el for el in d.find_elements(By.XPATH, f"//*[contains(text(), {grade_token!r}) and contains(text(), {letter!r})]") if el.is_displayed()), None))
        except TimeoutException:
            raise RuntimeError(f"nÃ£o encontrei a opÃ§Ã£o para a turma {target_class} na tela")

    # Garante que o elemento estÃ¡ visÃ­vel e clica via JS para evitar interceptaÃ§Ã£o
    driver.execute_script("arguments[0].scrollIntoView({block: 'center'});", option)
    time.sleep(0.5)
    driver.execute_script("arguments[0].click()", option)
        
    time.sleep(1.0)
    print("Confirmando selecao da turma com clique direto...")
    try:
        click_text(driver, "Selecionar turma", 5)
    except Exception:
        click_submit_like(driver, "Selecionar turma", 5)

    # Validacao final: espera a pagina de sucesso/material carregar
    time.sleep(2)
    print("Finalizando processo de selecao de turma.")


def click_if_present(driver, text, seconds=3):
    try:
        click_text(driver, text, seconds)
        return True
    except TimeoutException:
        return False


def visible_inputs(driver):
    return [field for field in driver.find_elements(By.CSS_SELECTOR, "input") if field.is_displayed()]


def visible_input_by_index(driver, index):
    inputs = driver.execute_script(
        """
        const norm = (value) => (value || '')
          .normalize('NFD').replace(/[\\u0300-\\u036f]/g, '')
          .replace(/[^a-zA-Z0-9\\s]/g, '')
          .replace(/\\s+/g, ' ').trim().toLowerCase();
        const username = [...document.querySelectorAll('input, textarea')]
          .find((el) => norm(el.getAttribute('placeholder')).includes('usu'));
        if (username && arguments[0] === 1) return [null, username];
        const visible = (el) => {
          const r = el.getBoundingClientRect();
          const st = getComputedStyle(el);
          return r.width > 0 && r.height > 0 && st.visibility !== 'hidden' && st.display !== 'none';
        };
        return [...document.querySelectorAll('input')]
          .filter(visible)
          .sort((a, b) => {
            const ar = a.getBoundingClientRect();
            const br = b.getBoundingClientRect();
            return (ar.top - br.top) || (ar.left - br.left);
          });
        """,
        index,
    )
    if len(inputs) <= index:
        raise RuntimeError(f"input visivel #{index + 1} nao encontrado; total={len(inputs)}")
    return inputs[index]


def field_hint(driver, field):
    attrs = []
    for attr in ("name", "type", "placeholder", "aria-label", "autocomplete", "value"):
        value = field.get_attribute(attr)
        if value:
            attrs.append(value)
    try:
        ancestor = driver.execute_script(
            """
            const el = arguments[0];
            const parts = [];
            let cur = el;
            for (let i = 0; i < 5 && cur; i++, cur = cur.parentElement) {
              if (cur.innerText) parts.push(cur.innerText);
              if (cur.previousElementSibling?.innerText) parts.push(cur.previousElementSibling.innerText);
            }
            return parts.join(' ');
            """,
            field,
        )
    except Exception:
        ancestor = ""
    return normalize_text(" ".join(attrs + [ancestor[:250]]))


def input_near_text(driver, text):
    return driver.execute_script(
        """
        const rawTarget = arguments[0];
        const norm = (value) => (value || '')
          .normalize('NFD').replace(/[\\u0300-\\u036f]/g, '')
          .replace(/[^a-zA-Z0-9\\s]/g, '')
          .replace(/\\s+/g, ' ').trim().toLowerCase();
        const targets = [norm(rawTarget)];
        if (targets[0].includes('usuario')) {
          targets.push('nome de usuario', 'nome de usu', 'usuario');
        }
        const visible = (el) => {
          const r = el.getBoundingClientRect();
          const st = getComputedStyle(el);
          return r.width > 0 && r.height > 0 && st.visibility !== 'hidden' && st.display !== 'none';
        };
        const inputs = [...document.querySelectorAll('input')].filter(visible);
        const labels = [...document.querySelectorAll('body *')]
          .filter((el) => visible(el) && targets.some((target) => norm(el.innerText || el.textContent).includes(target)))
          .map((el) => el.getBoundingClientRect())
          .filter((r) => r.width > 0 && r.height > 0);
        let best = null;
        let bestScore = Infinity;
        for (const input of inputs) {
          const r = input.getBoundingClientRect();
          const meta = norm([input.name, input.type, input.placeholder, input.getAttribute('aria-label'), input.autocomplete].join(' '));
          if (targets.some((target) => meta.includes(target))) return input;
          for (const label of labels) {
            const below = r.top >= label.top - 8;
            const close = Math.abs(r.top - label.top) < 90 || Math.abs(r.top - label.bottom) < 90;
            const horizontal = r.left < label.right + 260 && r.right > label.left - 260;
            if (below && close && horizontal) {
              const score = Math.abs(r.top - label.bottom) + Math.abs(r.left - label.left) / 10;
              if (score < bestScore) {
                bestScore = score;
                best = input;
              }
            }
          }
        }
        return best;
        """,
        text,
    )


def password_fields_on_page(driver):
    matches = []
    for field in visible_inputs(driver):
        hint = field_hint(driver, field)
        field_type = normalize_text(field.get_attribute("type"))
        if field_type == "password" or "senha" in hint or "password" in hint:
            if field not in matches:
                matches.append(field)
    return matches


def fill_signup_passwords(driver):
    try:
        pw_fields = wait(driver, 10).until(lambda d: password_fields_on_page(d))
    except TimeoutException:
        if email_exists_on_page(driver):
            raise EmailAlreadyExists("email ja existe no Geekie One; use a conta existente")
        result = driver.execute_script(
            """
            const password = arguments[0];
            const norm = (value) => (value || '')
              .normalize('NFD').replace(/[\\u0300-\\u036f]/g, '')
              .replace(/[^a-zA-Z0-9\\s]/g, '')
              .replace(/\\s+/g, ' ').trim().toLowerCase();
            const fields = [...document.querySelectorAll('input')]
              .filter((el) => (el.type || '').toLowerCase() === 'password'
                || norm(el.placeholder || el.getAttribute('aria-label') || '').includes('senha'));
            const setter = Object.getOwnPropertyDescriptor(HTMLInputElement.prototype, 'value')?.set;
            for (const field of fields) {
              field.focus();
              field.click();
              if (setter) setter.call(field, password);
              else field.value = password;
              field.dispatchEvent(new InputEvent('input', {bubbles: true, inputType: 'insertText', data: password}));
              field.dispatchEvent(new Event('change', {bubbles: true}));
              field.dispatchEvent(new Event('blur', {bubbles: true}));
            }
            return fields.length;
            """,
            PASSWORD,
        )
        if result:
            return []
        raise
    for _ in range(3):
        for pw_field in pw_fields:
            fill_field(driver, pw_field, PASSWORD)
            try:
                pw_field.send_keys(Keys.TAB)
            except Exception:
                pass
        time.sleep(0.4)
        empty = [field for field in pw_fields if not (field.get_attribute("value") or "")]
        if not empty:
            return pw_fields
        pw_fields = password_fields_on_page(driver)
    lengths = [len(field.get_attribute("value") or "") for field in pw_fields]
    raise RuntimeError(f"senha nao preenchida no cadastro (campos={len(pw_fields)}, tamanhos={lengths})")


def input_by_placeholder(driver, text, mode="contains"):
    found = driver.execute_script(
        """
        const rawTarget = arguments[0];
        const mode = arguments[1];
        const norm = (value) => (value || '')
          .normalize('NFD').replace(/[\\u0300-\\u036f]/g, '')
          .replace(/[^a-zA-Z0-9\\s]/g, '')
          .replace(/\\s+/g, ' ').trim().toLowerCase();
        const target = norm(rawTarget);
        const inputs = [...document.querySelectorAll('input, textarea')];
        for (const input of inputs) {
          const placeholder = norm(input.getAttribute('placeholder'));
          const name = norm(input.getAttribute('name'));
          const aria = norm(input.getAttribute('aria-label'));
          const meta = [placeholder, name, aria].join(' ');
          if ((mode === 'starts' && placeholder.startsWith(target)) || meta.includes(target)) {
            return input;
          }
        }
        if (target.includes('usuario')) {
          return inputs.find((input) => norm(input.getAttribute('placeholder')).includes('usu')) || null;
        }
        return null;
        """,
        text,
        mode,
    )
    if found:
        return found
    if mode == "starts":
        xpath = f"//input[starts-with(@placeholder, {text!r})]"
    else:
        xpath = f"//input[contains(@placeholder, {text!r})]"
    matches = driver.find_elements(By.XPATH, xpath)
    visible = [field for field in matches if field.is_displayed()]
    if not visible:
        raise RuntimeError(f"campo nao encontrado: {text}")
    return visible[0]


def input_by_hint(driver, text):
    target = normalize_text(normalize_ascii(text))
    alternatives = {target}
    if "usuario" in target:
        alternatives.update({"nome de usu", "usuario"})
    near = input_near_text(driver, text)
    if near:
        return near
    matches = []
    for field in visible_inputs(driver):
        hint = field_hint(driver, field)
        if any(item in hint for item in alternatives):
            matches.append(field)
    if matches:
        return matches[0]
    if "usuario" in target:
        inputs = driver.execute_script(
            """
            const visible = (el) => {
              const r = el.getBoundingClientRect();
              const st = getComputedStyle(el);
              return r.width > 0 && r.height > 0 && st.visibility !== 'hidden' && st.display !== 'none';
            };
            return [...document.querySelectorAll('input')]
              .filter(visible)
              .sort((a, b) => {
                const ar = a.getBoundingClientRect();
                const br = b.getBoundingClientRect();
                return (ar.top - br.top) || (ar.left - br.left);
              });
            """
        )
        if len(inputs) >= 2:
            return inputs[1]
    return input_by_placeholder(driver, text)


def fill_signup_form(driver, name, email, username):
    wait(driver, 20).until(lambda d: d.execute_script("return document.querySelectorAll('input, textarea').length") >= 3)
    dismiss_notification_popup(driver)
    filled = driver.execute_script(
        """
        const values = arguments[0];
        const norm = (value) => (value || '')
          .normalize('NFD').replace(/[\\u0300-\\u036f]/g, '')
          .replace(/[^a-zA-Z0-9\\s]/g, '')
          .replace(/\\s+/g, ' ').trim().toLowerCase();
        const inputs = [...document.querySelectorAll('input, textarea')];
        const setValue = (el, value) => {
          el.focus();
          el.click();
          const proto = el.tagName === 'TEXTAREA' ? HTMLTextAreaElement.prototype : HTMLInputElement.prototype;
          const setter = Object.getOwnPropertyDescriptor(proto, 'value')?.set;
          if (setter) setter.call(el, value);
          else el.value = value;
          el.dispatchEvent(new InputEvent('input', {bubbles: true, inputType: 'insertText', data: value}));
          el.dispatchEvent(new Event('change', {bubbles: true}));
          el.dispatchEvent(new Event('blur', {bubbles: true}));
        };
        const ph = (el) => norm(el.getAttribute('placeholder'));
        const nameInput = inputs.find((el) => ph(el).includes('nome completo'));
        const emailInputs = inputs.filter((el) => ph(el).includes('mail') && ph(el).includes('estudante'));
        const emailInput = emailInputs[0] || inputs.find((el) => ph(el).includes('email') || ph(el).includes('e mail'));
        const confirmInput = inputs.find((el) => ph(el).includes('confirme')) || emailInputs[1];
        const userInput = inputs.find((el) => ph(el).includes('usu'));
        if (!nameInput || !emailInput || !confirmInput || !userInput) {
          return {ok: false, placeholders: inputs.map((el) => el.getAttribute('placeholder') || '')};
        }
        setValue(nameInput, values.name);
        setValue(emailInput, values.email);
        setValue(confirmInput, values.email);
        setValue(userInput, values.username);
        return {ok: true};
        """,
        {"name": name, "email": email, "username": username},
    )
    if not filled or not filled.get("ok"):
        raise RuntimeError(f"campos basicos do cadastro nao encontrados: {filled}")
    click_if_present(driver, "CHECAR", 2)

def username_exists_on_page(driver):
    page = normalize_text(normalize_ascii(driver.page_source))
    return (
        "nome de usuario ja existe" in page
        or "usuario ja existe" in page
        or "nome de usuario ja existente" in page
        or "usuario ja existente" in page
        or "usuario em uso" in page
    )


def email_exists_on_page(driver):
    page = normalize_text(normalize_ascii(driver.page_source))
    return (
        "e-mail ja esta em uso" in page
        or "email ja esta em uso" in page
        or "mail ja esta em uso" in page
        or "acesse a conta existente" in page
    )


def wait_email_exists_on_page(driver, seconds=4):
    deadline = time.time() + seconds
    while time.time() < deadline:
        if email_exists_on_page(driver):
            return True
        time.sleep(0.4)
    return False


def ensure_available_signup_username(driver, username):
    current = username
    for number in [""] + [str(item) for item in range(2, 100)]:
        candidate = username if not number else f"{username}{number}"
        fill_field(driver, input_by_hint(driver, "nome de usuario"), candidate)
        click_if_present(driver, "CHECAR", 2)
        time.sleep(0.8)
        if not username_exists_on_page(driver):
            return candidate
        current = candidate
    raise RuntimeError(f"nome de usuario indisponivel: {current}")


def click_enabled_text(driver, text, seconds=10):
    deadline = time.time() + seconds
    xpath = f"//*[self::button or @role='button' or self::a][.//*[normalize-space()={text!r}] or normalize-space()={text!r}]"
    while time.time() < deadline:
        for element in reversed(driver.find_elements(By.XPATH, xpath)):
            disabled = element.get_attribute("disabled")
            aria_disabled = element.get_attribute("aria-disabled")
            if element.is_displayed() and not disabled and aria_disabled != "true":
                driver.execute_script("arguments[0].click()", element)
                return True
        time.sleep(0.25)
    return False


def click_submit_like(driver, texts, seconds=15):
    if isinstance(texts, str):
        texts = [texts]
    normalized_targets = [normalize_text(text) for text in texts]
    deadline = time.time() + seconds
    while time.time() < deadline:
        try:
            clicked = driver.execute_script(
                """
            const targets = arguments[0];
            const norm = (value) => (value || '')
              .normalize('NFD').replace(/[\\u0300-\\u036f]/g, '')
              .replace(/[^a-zA-Z0-9\\s]/g, '')
              .replace(/\\s+/g, ' ').trim().toLowerCase();
            const visible = (el) => {
              const r = el.getBoundingClientRect();
              const st = getComputedStyle(el);
              return r.width > 0 && r.height > 0 && st.visibility !== 'hidden' && st.display !== 'none';
            };
            const elements = [...document.querySelectorAll('button, [role=button], input[type=submit], a, div, span')]
              .filter(visible)
              .filter((el) => {
                const text = norm(el.innerText || el.textContent || el.value || el.getAttribute('aria-label'));
                return targets.some((target) => text.includes(target));
              });
            for (const el of elements.reverse()) {
              const disabled = el.disabled || el.getAttribute('aria-disabled') === 'true';
              const rawClass = el.className && (el.className.baseVal || el.className.toString?.() || el.className);
              const cls = norm(rawClass);
              if (!disabled && !cls.includes('disabled')) {
                el.scrollIntoView({block: 'center'});
                el.click();
                return true;
              }
            }
            return false;
            """,
                normalized_targets,
            )
        except Exception:
            clicked = False
        if clicked:
            time.sleep(1)
            return True
        time.sleep(0.5)
    return False


def click_continue_anywhere(driver, seconds=12):
    deadline = time.time() + seconds
    while time.time() < deadline:
        try:
            clicked = driver.execute_script(
                """
                const norm = (value) => (value || '')
                  .normalize('NFD').replace(/[\\u0300-\\u036f]/g, '')
                  .replace(/[^a-zA-Z0-9\\s]/g, '')
                  .replace(/\\s+/g, ' ').trim().toLowerCase();
                const nodes = [...document.querySelectorAll('[role=button], button, a, div, span')]
                  .filter((el) => norm(el.innerText || el.textContent || el.getAttribute('aria-label')).includes('continuar'));
                for (const node of nodes.reverse()) {
                  const target = node.closest('[role=button]') || node.closest('button') || node.closest('a') || node;
                  const disabled = target.disabled || target.getAttribute('aria-disabled') === 'true';
                  if (disabled) continue;
                  target.scrollIntoView({block: 'center'});
                  for (const type of ['pointerdown', 'mousedown', 'mouseup', 'click']) {
                    target.dispatchEvent(new MouseEvent(type, {bubbles: true, cancelable: true, view: window}));
                  }
                  target.click();
                  return true;
                }
                return false;
                """
            )
        except Exception:
            clicked = False
        if clicked:
            time.sleep(0.8)
            return True
        time.sleep(0.4)
    return False


def click_continue_by_coordinates(driver):
    size = driver.execute_script("return {w: window.innerWidth, h: window.innerHeight};")
    x = int(size["w"] * 0.50)
    for y_factor in (0.74, 0.68, 0.80, 0.60):
        cdp_click(driver, x, int(size["h"] * y_factor))
        time.sleep(1)
        if driver.find_elements(By.XPATH, "//input[contains(@placeholder, 'nome completo')]"):
            return True
    return False


def click_continue_by_dom_rect(driver):
    rect = driver.execute_script(
        """
        const norm = (value) => (value || '')
          .normalize('NFD').replace(/[\\u0300-\\u036f]/g, '')
          .replace(/[^a-zA-Z0-9\\s]/g, '')
          .replace(/\\s+/g, ' ').trim().toLowerCase();
        const nodes = [...document.querySelectorAll('[role=button], button, a, div, span')]
          .filter((el) => norm(el.innerText || el.textContent || el.getAttribute('aria-label')).includes('continuar'));
        for (const node of nodes.reverse()) {
          const target = node.closest('[role=button]') || node.closest('button') || node.closest('a') || node;
          const r = target.getBoundingClientRect();
          if (r.width > 0 && r.height > 0) {
            target.scrollIntoView({block: 'center'});
            const rr = target.getBoundingClientRect();
            return {x: rr.left + rr.width / 2, y: rr.top + rr.height / 2};
          }
        }
        return null;
        """
    )
    if not rect:
        return False
    cdp_click(driver, int(rect["x"]), int(rect["y"]))
    time.sleep(1)
    return bool(driver.find_elements(By.XPATH, "//input[contains(@placeholder, 'nome completo')]"))


def continue_parent_warning(driver):
    for attempt in range(4):
        try:
            click_continue_anywhere(driver, 4)
            wait(driver, 4).until(
                lambda d: d.find_elements(By.XPATH, "//input[contains(@placeholder, 'nome completo')]")
            )
            return True
        except Exception as exc:
            print(f"Aviso: clique JS no continuar falhou/tela nao mudou: {str(exc).splitlines()[0] or type(exc).__name__}")
        try:
            if click_continue_by_dom_rect(driver):
                return True
        except Exception as exc:
            print(f"Aviso: clique por retangulo no continuar falhou: {str(exc).splitlines()[0] or type(exc).__name__}")
        try:
            if click_continue_by_coordinates(driver):
                return True
        except Exception as exc:
            print(f"Aviso: clique por coordenada no continuar falhou: {str(exc).splitlines()[0] or type(exc).__name__}")
        print(f"Aviso: tentativa {attempt + 1} de continuar no aviso do responsavel nao abriu cadastro.")
    return False


def select_created_account_class_safe(driver, class_code):
    select = wait(driver, 20).until(EC.element_to_be_clickable((By.XPATH, "//input[@readonly and (contains(@value, 'turma') or contains(@class, 'input'))]")))
    grade_token, letter, target_class = class_parts(class_code)
    
    # Tenta abrir o menu de vÃ¡rias formas
    opened = False
    for method in ["Overlay click", "JS click", "Native click", "Space key", "Container click"]:
        try:
            if method == "Overlay click":
                overlays = driver.find_elements(By.CSS_SELECTOR, "[data-testid='DropdownInput-overlay']")
                if overlays:
                    driver.execute_script("arguments[0].click()", overlays[0])
                else:
                    continue
            elif method == "JS click":
                driver.execute_script("arguments[0].click()", select)
            elif method == "Native click":
                select.click()
            elif method == "Space key":
                select.send_keys(Keys.SPACE)
            elif method == "Container click":
                parent = select.find_element(By.XPATH, "./..")
                driver.execute_script("arguments[0].click()", parent)
            
            time.sleep(2)
            if driver.find_elements(By.XPATH, f"//*[(@role='option' or @role='button' or self::li) and contains(., {grade_token!r})]"):
                opened = True
                break
        except Exception:
            continue
            
    print(f"Selecionando turma (cadastro novo): {target_class}")
    
    option_selectors = [
        f"//*[(@role='option' or @role='button') and contains(., {grade_token!r}) and contains(., {letter!r})]",
        f"//li[contains(., {grade_token!r}) and contains(., {letter!r})]",
        f"//div[contains(., {grade_token!r}) and contains(., {letter!r}) and not(contains(@class, 'input'))]"
    ]
    
    option = None
    for selector in option_selectors:
        try:
            elements = driver.find_elements(By.XPATH, selector)
            visible = [el for el in elements if el.is_displayed() and el.size['width'] > 0]
            if visible:
                option = visible[0]
                break
        except Exception:
            continue
            
    if not option:
        try:
            option = wait(driver, 10).until(lambda d: next((el for el in d.find_elements(By.XPATH, f"//*[contains(text(), {grade_token!r}) and contains(text(), {letter!r})]") if el.is_displayed()), None))
        except TimeoutException:
            raise RuntimeError(f"nÃ£o encontrei a opÃ§Ã£o para a turma {target_class} na tela")

    driver.execute_script("arguments[0].scrollIntoView({block: 'center'});", option)
    time.sleep(0.5)
    driver.execute_script("arguments[0].click()", option)
        
    time.sleep(1.0)
    print("Confirmando selecao da turma com clique direto...")
    try:
        click_text(driver, "Selecionar turma", 5)
    except Exception:
        click_submit_like(driver, "Selecionar turma", 5)

    # Validacao final: espera a pagina de sucesso/material carregar
    time.sleep(2)
    print("Finalizando processo de selecao de turma.")


def create_one_account(driver, name, email, code, class_code, used_usernames):
    # O e-mail usado DEVE ser o da planilha
    print(f"Criando conta nova para {name} com e-mail: {email}")
    username = username_from_email(email, used_usernames)
    ensure_one_logged_out(driver)
    driver.get(ONE_URL)
    dismiss_notification_popup(driver)
    click_text(driver, "Ativar novo material")
    field = wait(driver).until(EC.visibility_of_element_located((By.CSS_SELECTOR, "input[type='text']")))
    field.send_keys(str(code))
    click_text(driver, "Ativar")
    click_text(driver, "Criar conta", 20)
    time.sleep(1)
    if not driver.find_elements(By.XPATH, "//input[contains(@placeholder, 'nome completo')]"):
        print(f"\n{name}: tela de aviso do responsavel aberta. Tentando continuar automaticamente...")
        try:
            Path(__file__).with_name("debug_responsavel_antes.html").write_text(driver.page_source, encoding="utf-8")
            driver.save_screenshot(str(Path(__file__).with_name("debug_responsavel_antes.png")))
        except Exception:
            pass
        if not continue_parent_warning(driver):
            try:
                Path(__file__).with_name("debug_responsavel_depois.html").write_text(driver.page_source, encoding="utf-8")
                driver.save_screenshot(str(Path(__file__).with_name("debug_responsavel_depois.png")))
            except Exception:
                pass
            input("Aviso: nao consegui clicar em CONTINUAR automaticamente. Clique no Chrome e pressione ENTER aqui...")
    print("Cadastro aberto; preenchendo formulario...")
    dismiss_notification_popup(driver)
    print("Cadastro: preenchendo nome/e-mail/usuario...")
    fill_signup_form(driver, name, email, username)
    print("Cadastro: formulario basico preenchido.")
    if wait_email_exists_on_page(driver, 5):
        raise EmailAlreadyExists("email ja existe no Geekie One; use a conta existente")
    print("Cadastro: checando usuario disponivel...")
    username = ensure_available_signup_username(driver, username)
    print(f"Cadastro: usuario definido como {username}.")
    if wait_email_exists_on_page(driver, 2):
        raise EmailAlreadyExists("email ja existe no Geekie One; use a conta existente")
    time.sleep(1)
    print("Cadastro: clicando em continuar na etapa de dados...")
    if not (
        click_continue_anywhere(driver, 15)
        or click_submit_like(driver, ["Continuar", "Criar conta", "Cadastrar"], 5)
        or click_enabled_text(driver, "Continuar", 3)
        or click_enabled_text(driver, "Criar conta", 3)
        or click_enabled_text(driver, "Cadastrar", 3)
    ):
        raise RuntimeError("nao encontrei botao para continuar dados do cadastro")
    print("Tela de dados enviada; preenchendo senha...")
    fill_signup_passwords(driver)
    print("Cadastro: senha preenchida.")
    time.sleep(1)
    print("Cadastro: finalizando cadastro...")
    if not (
        click_submit_like(driver, ["Salvar alteraÃ§Ãµes", "Salvar alteracoes", "Continuar", "Criar conta", "Cadastrar"], 15)
        or click_enabled_text(driver, "Continuar", 3)
        or click_enabled_text(driver, "Criar conta", 3)
        or click_enabled_text(driver, "Cadastrar", 3)
    ):
        raise RuntimeError("nao encontrei botao para finalizar cadastro")
    time.sleep(2)
    try:
        updated_username = complete_profile_after_login(driver, email, used_usernames)
        if updated_username:
            username = updated_username
    except Exception as exc:
        detail = str(exc).splitlines()[0] or type(exc).__name__
        print(f"Aviso: atualizacao pos-cadastro falhou pelo fluxo completo: {detail}")
        select_no_cpf(driver)
        if not (click_save_profile_by_cdp_dom(driver) or click_save_profile(driver)):
            raise RuntimeError("nao consegui salvar a atualizacao das informacoes pos-cadastro")
    for text in ("Pular", "Agora nÃ£o", "NÃ£o quero informar", "Nao quero informar"):
        click_if_present(driver, text, 2)
    click_if_present(driver, "Continuar", 5)
    click_if_present(driver, "Acessar seu material Geekie One!", 15)
    select_created_account_class_safe(driver, class_code)
    return username


def logout(driver):
    try:
        profile = wait(driver, 8).until(EC.element_to_be_clickable((By.CSS_SELECTOR, "[aria-label='Perfil']")))
        driver.execute_script("arguments[0].click()", profile)
        click_text(driver, "Sair", 8)
    except TimeoutException:
        pass


def resolve_sheet(workbook, requested_name, sheet_index=None):
    if sheet_index is not None:
        try:
            return workbook.worksheets[sheet_index - 1]
        except IndexError:
            raise RuntimeError(f"Aba de indice {sheet_index} nao existe. Abas: {workbook.sheetnames}")

    if normalize_text(requested_name) in {"", "auto"}:
        geekie_sheets = [name for name in workbook.sheetnames if "geekie" in normalize_text(name)]
        if len(geekie_sheets) == 1:
            return workbook[geekie_sheets[0]]
        available = ", ".join(workbook.sheetnames)
        raise KeyError(f"Nao foi possivel escolher a aba automaticamente. Abas disponiveis: {available}")

    if requested_name in workbook.sheetnames:
        return workbook[requested_name]

    normalized_requested = normalize_text(requested_name)
    normalized_names = {name: normalize_text(name) for name in workbook.sheetnames}

    for name, normalized in normalized_names.items():
        if normalized == normalized_requested:
            return workbook[name]

    for name, normalized in normalized_names.items():
        if normalized_requested and normalized_requested in normalized:
            return workbook[name]

    geekie_sheets = [name for name, normalized in normalized_names.items() if "geekie" in normalized]
    if len(geekie_sheets) == 1:
        return workbook[geekie_sheets[0]]

    available = ", ".join(workbook.sheetnames)
    raise KeyError(
        f"A aba {requested_name!r} nao existe. Abas disponiveis: {available}"
    )


def detect_layout(worksheet):
    header_map = {}
    for cell in worksheet[1]:
        normalized = normalize_text(cell.value)
        if normalized:
            header_map[normalized] = cell.column_letter

    compact_layout = {
        "name": header_map.get("nomealuno"),
        "email": header_map.get("loginoffice365"),
        "username": header_map.get("usuario"),
        "code": header_map.get("codigo do material") or header_map.get("login / codigo do material"),
        "class_code": header_map.get("codigoturma"),
        "status": header_map.get("status"),
    }
    if all(compact_layout[key] for key in ("name", "email", "username", "code", "class_code")):
        return 2, compact_layout

    legacy_sample = {
        "name": worksheet["H2"].value,
        "email": worksheet["I2"].value,
        "username": worksheet["J2"].value,
        "code": worksheet["K2"].value,
        "class_code": worksheet["D2"].value,
    }
    legacy_layout = {
        "name": "H",
        "email": "I",
        "username": "J",
        "code": "K",
        "class_code": "D",
        "status": None,
    }
    if (
        legacy_sample["name"]
        and "@" in str(legacy_sample["email"] or "")
        and legacy_sample["code"]
        and legacy_sample["class_code"]
    ):
        return 2, legacy_layout

    raise RuntimeError(
        f"Formato da aba {worksheet.title!r} nao reconhecido para automacao."
    )


def search_lab(driver, name, email=None):
    driver.get(LAB_URL)
    wait(driver, 12).until(lambda d: "educadores.geekieteste.geekie.com.br" in d.current_url)
    current = driver.current_url
    if "/application_report/" in current:
        users_url = current.split("/application_report/", 1)[0] + "/manage/users"
    elif "/manage/" in current:
        users_url = current.split("/manage/", 1)[0] + "/manage/users"
    else:
        users_url = current.rstrip("/") + "/manage/users"
    driver.get(users_url)

    search = wait(driver).until(EC.element_to_be_clickable((
        By.XPATH,
        "//input[contains(@placeholder, 'Buscar') or @type='search']",
    )))
    target_name = normalize_text(name)
    name_parts = [part for part in normalize_ascii(name).split() if part]

    def lab_search_terms():
        terms = [str(name)]
        if len(name_parts) >= 3:
            terms.append(" ".join(name_parts[:3]))
        if len(name_parts) >= 2:
            terms.append(" ".join(name_parts[:2]))
            terms.append(f"{name_parts[0]} {name_parts[-1]}")
        if name_parts:
            terms.append(name_parts[0])
        if email:
            local = str(email).split("@", 1)[0]
            local = re.sub(r"-escola.*$", "", local)
            email_parts = [part for part in re.split(r"[._\\-]+", local) if part]
            if len(email_parts) >= 2:
                terms.append(" ".join(email_parts[:2]))
            if email_parts:
                terms.append(email_parts[0])
        seen = set()
        unique = []
        for term in terms:
            key = normalize_text(term)
            if key and key not in seen:
                seen.add(key)
                unique.append(term)
        return unique

    def type_search(query):
        current_search = wait(driver).until(EC.element_to_be_clickable((
            By.XPATH,
            "//input[contains(@placeholder, 'Buscar') or @type='search']",
        )))
        current_search.clear()
        current_search.send_keys(Keys.CONTROL, "a")
        current_search.send_keys(str(query))
        time.sleep(1)

    def include_removed_users():
        print("Ativando filtro para incluir removidos no Lab...")
        clicked = driver.execute_script(
            """
            const norm = (value) => (value || '')
              .normalize('NFD').replace(/[\\u0300-\\u036f]/g, '')
              .replace(/[^a-zA-Z0-9\\s]/g, '')
              .replace(/\\s+/g, ' ').trim().toLowerCase();
            const visible = (el) => {
              const r = el.getBoundingClientRect();
              const st = getComputedStyle(el);
              return r.width > 0 && r.height > 0 && st.visibility !== 'hidden' && st.display !== 'none';
            };
            const nodes = [...document.querySelectorAll('button, [role=button], label, md-checkbox, input, span, div')]
              .filter((el) => visible(el) && norm([el.innerText, el.textContent, el.getAttribute('aria-label'), el.title, el.value].join(' ')).includes('removidos'));
            for (const node of nodes) {
              const target = node.closest('label') || node.closest('[role=button]') || node.closest('md-checkbox') || node;
              target.click();
              target.dispatchEvent(new MouseEvent('click', {bubbles: true, cancelable: true, view: window}));
              return true;
            }
            return false;
            """
        )
        if not clicked:
            raise AccountNotFound("conta nao encontrada no Lab")
        time.sleep(2)

    def matching_name(d, query):
        query_norm = normalize_text(query)
        expected_email = normalize_text(email).replace(" ", "") if email else ""
        # Primeiro, busca por todos os cards de usuÃ¡rio visÃ­veis
        cards = d.find_elements(By.CSS_SELECTOR, ".user-card")
        visible_cards = [card for card in cards if card.is_displayed()]
        if len(visible_cards) == 1:
            # Se houver apenas um card, assume que Ã© o aluno (confia na busca do Lab)
            print(f"Aviso: apenas um resultado encontrado. Assumindo que Ã© o aluno correto.")
            return visible_cards[0]

        matches = []
        # Se houver mais de um, tenta encontrar o match exato (normalizado)
        for element in visible_cards:
            try:
                element_text = normalize_text(element.text)
                compact_text = element_text.replace(" ", "")
                if expected_email and expected_email in compact_text:
                    matches.append(element)
                    continue

                name_line = ""
                for line in element.text.splitlines():
                    normalized_line = normalize_text(line)
                    if normalized_line and "estudante" not in normalized_line and "colegio" not in normalized_line:
                        name_line = normalized_line
                        break

                card_words = name_line.split()
                target_words = target_name.split()
                prefix_match = (
                    len(card_words) >= 2
                    and len(target_words) >= len(card_words)
                    and target_words[: len(card_words)] == card_words
                )
                exact_match = name_line == target_name
                if exact_match or prefix_match:
                    matches.append(element)
            except StaleElementReferenceException:
                continue
        if not matches:
            return False

        def is_inactive(card):
            """Detects both the old CSS marker and the visible Lab status text."""
            try:
                if card.find_elements(By.CSS_SELECTOR, ".inactive-warning:not(.ng-hide)"):
                    return True
                card_text = normalize_text(card.text)
                return "usuario inativo" in card_text or "usuaria inativa" in card_text
            except StaleElementReferenceException:
                return True

        # A search can return an old removed account and the current account
        # with the same name. The removed card often has no reset-link action.
        # Keep it only as a fallback when no active match exists.
        active_matches = [element for element in matches if not is_inactive(element)]
        if active_matches:
            if len(active_matches) != len(matches):
                print("Aviso: resultados duplicados no Lab; priorizando o usuario ativo.")
            matches = active_matches

        def has_info_action(card):
            try:
                buttons = card.find_elements(
                    By.XPATH,
                    ".//button[contains(@tooltip, 'Ver') and "
                    "(contains(@tooltip, 'informa') or contains(@tooltip, 'Informa'))]",
                )
                return bool(buttons or card.find_elements(By.CSS_SELECTOR, "button[tooltip]"))
            except StaleElementReferenceException:
                return False

        try:
            # The Lab places the active duplicate below the removed one. When
            # both records expose the same e-mail and the status marker is not
            # available in the card text, prefer the later actionable card.
            ranked = sorted(
                enumerate(matches),
                key=lambda item: (is_inactive(item[1]), not has_info_action(item[1]), -item[0]),
            )
            return ranked[0][1]
        except StaleElementReferenceException:
            return False

    def find_student():
        # Busca pelo nome (conforme solicitado, o Lab nÃ£o busca por e-mail)
        last_error = None
        for term in lab_search_terms():
            print(f"Buscando no Lab: {term}")
            type_search(term)
            try:
                return wait(driver, 5).until(lambda d, q=term: matching_name(d, q))
            except TimeoutException as exc:
                last_error = exc
                continue
        raise last_error or TimeoutException("conta nao encontrada no Lab")

    try:
        student_name = find_student()
    except TimeoutException:
        print("Aluno nao encontrado sem removidos; repetindo busca com removidos incluidos.")
        include_removed_users()
        try:
            student_name = find_student()
            print("Aviso: aluno encontrado no Lab ao incluir removidos.")
        except TimeoutException:
            raise AccountNotFound("conta nao encontrada no Lab")

    try:
        card = student_name.find_element(By.XPATH, "./ancestor-or-self::*[contains(@class, 'user-card')][1]")
    except Exception:
        visible_cards = [card for card in driver.find_elements(By.CSS_SELECTOR, ".user-card") if card.is_displayed()]
        if len(visible_cards) == 1:
            card = visible_cards[0]
        else:
            raise

    # Verificar se o aluno estÃ¡ inativo â€” nÃ£o bloquear, apenas avisar e continuar
    inactive = card.find_elements(By.CSS_SELECTOR, ".inactive-warning:not(.ng-hide)")
    if not inactive:
        try:
            inactive = [card] if "usuario inativo" in normalize_text(card.text) else []
        except StaleElementReferenceException:
            inactive = []
    if inactive:
        print(f"Aviso: usuario inativo no Lab; tentando obter link de reset mesmo assim.")

    info_buttons = card.find_elements(By.XPATH, ".//button[contains(@tooltip, 'Ver') and (contains(@tooltip, 'informa') or contains(@tooltip, 'Informa'))]")
    if not info_buttons:
        info_buttons = card.find_elements(By.CSS_SELECTOR, "button[tooltip]")
    if not info_buttons:
        raise RuntimeError("botao de informacoes nao encontrado no Lab")
    driver.execute_script("arguments[0].click()", info_buttons[0])

    wait(driver).until(EC.visibility_of_element_located((By.XPATH, "//*[contains(., 'Dados de acesso')]")))
    readonly = driver.find_elements(By.CSS_SELECTOR, "input[readonly]")
    values = [item.get_attribute("value") for item in readonly if item.get_attribute("value")]
    if not values:
        raise RuntimeError("usuario nao encontrado nos dados de acesso")
    username = values[-1]
    
    # CONFERÃŠNCIA DE E-MAIL: Verifica se algum dos valores encontrados Ã© o e-mail da planilha
    found_emails = [v.lower().strip() for v in values if "@" in v]
    if email and email.lower().strip() not in found_emails:
        print(f"Aviso: Aluno encontrado no Lab tem e-mail diferente ({found_emails}). Ignorando para criar conta nova.")
        raise AccountNotFound("e-mail divergente no Lab")
        
    click_text(driver, "Obter link")
    link = wait(driver).until(
        EC.presence_of_element_located((By.CSS_SELECTOR, "a[href*='reset_password_token=']"))
    ).get_attribute("href")
    return username, link


def run(args):
    workbook_path = Path(args.workbook).resolve()
    wb = load_workbook(workbook_path)
    ws = resolve_sheet(wb, args.sheet, args.sheet_index)
    first_data_row, columns = detect_layout(ws)
    driver = new_driver()
    try:
        ensure_logins(driver)
        start_row = max(args.start, first_data_row)
        used_usernames = [
            ws[f"{columns['username']}{candidate_row}"].value
            for candidate_row in range(first_data_row, ws.max_row + 1)
            if ws[f"{columns['username']}{candidate_row}"].value
        ]
        for row in range(start_row, ws.max_row + 1):
            name = ws[f"{columns['name']}{row}"].value
            email = ws[f"{columns['email']}{row}"].value
            existing_username = ws[f"{columns['username']}{row}"].value
            code = ws[f"{columns['code']}{row}"].value
            class_code = ws[f"{columns['class_code']}{row}"].value
            status = ws[f"{columns['status']}{row}"].value if columns["status"] else None
            if not name or not email or existing_username:
                continue
            if not code:
                print(f"[{row}] {name}: sem codigo do material; pulando.")
                if columns["status"]:
                    ws[f"{columns['status']}{row}"] = "sem codigo"
                    wb.save(workbook_path)
                continue
            if normalize_text(status) == "ok":
                continue
            print(f"[{row}] {name}")
            # Guardar linha atual para screenshots de depuraÃ§Ã£o
            driver.current_row_num = row
            stage = "buscar aluno no Geekie Lab"
            try:
                try:
                    username, reset_link = search_lab(driver, str(name), str(email))
                    stage = "redefinir senha"
                    reset_password(driver, reset_link, str(name))
                    stage = "entrar no Geekie One"
                    # LOGIN POR E-MAIL (mais seguro)
                    updated_username = login_one(driver, str(email), used_usernames)
                    if updated_username:
                        username = updated_username
                    stage = "ativar material e selecionar turma"
                    activate(driver, str(code), str(class_code))
                except AccountNotFound:
                    stage = "criar conta no Geekie One"
                    username = create_one_account(
                        driver,
                        str(name),
                        str(email),
                        str(code),
                        str(class_code),
                        used_usernames,
                    )
                ws[f"{columns['username']}{row}"] = username
                if columns["status"]:
                    ws[f"{columns['status']}{row}"] = "ok"
                wb.save(workbook_path)
                used_usernames.append(username)
                print(f"Planilha atualizada: linha {row}, usuario {username}")
                stage = "sair do Geekie One"
                logout(driver)
                print(f"OK: {username}")
            except Exception as exc:
                lines = [line.strip() for line in str(exc).splitlines() if line.strip()]
                detail = lines[0] if lines else type(exc).__name__
                if detail.lower() in {"message:", "message"}:
                    detail = type(exc).__name__
                reason = f"{stage}: {detail} | URL: {driver.current_url}"
                log_error(row, str(name), reason)
                try:
                    driver.save_screenshot(str(Path(__file__).with_name(f"erro_linha_{row}.png")))
                    Path(__file__).with_name(f"erro_linha_{row}.html").write_text(
                        driver.page_source, encoding="utf-8"
                    )
                except Exception:
                    pass
                print(f"ERRO em {reason}")
                if not args.continue_on_error:
                    raise
    finally:
        driver.quit()


def main():
    parser = argparse.ArgumentParser(description="AutomaÃ§Ã£o local Geekie")
    sub = parser.add_subparsers(dest="command", required=True)
    sub.add_parser("clone-profile", help="copia a sessÃ£o do Chrome jÃ¡ usada")
    sub.add_parser("setup", help="cria o perfil e permite fazer login")
    process = sub.add_parser("run", help="processa uma planilha local")
    process.add_argument("workbook")
    process.add_argument("--sheet-index", type=int, default=5, help="Indice da aba, comecando em 1. Padrao: 5")
    process.add_argument("--sheet", default="9Âº Anos - Geekie")
    process.add_argument("--start", type=int, default=2)
    process.add_argument("--continue-on-error", action="store_true")
    args = parser.parse_args()
    if args.command == "clone-profile":
        clone_chrome_profile()
    elif args.command == "setup":
        setup_profile()
    else:
        run(args)


if __name__ == "__main__":
    try:
        main()
    except KeyboardInterrupt:
        sys.exit(130)
