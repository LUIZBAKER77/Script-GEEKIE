from __future__ import annotations

import argparse
import csv
import time
from pathlib import Path

from openpyxl import load_workbook
from selenium.webdriver.common.by import By

import auditoria_geekie as audit
import geekie_automation as g


LOG_FILE = Path(__file__).with_name("correcao_pendentes_log.csv")


def norm(value):
    return g.normalize_text(g.normalize_ascii("" if value is None else value))


def append_log(row):
    exists = LOG_FILE.exists()
    with LOG_FILE.open("a", newline="", encoding="utf-8") as file:
        writer = csv.writer(file)
        if not exists:
            writer.writerow(["aba", "linha", "nome", "resultado", "detalhe"])
        writer.writerow(row)


def completed_from_log():
    completed = set()
    if not LOG_FILE.exists():
        return completed
    with LOG_FILE.open("r", newline="", encoding="utf-8") as file:
        reader = csv.DictReader(file)
        for row in reader:
            if norm(row.get("resultado")) == "ok":
                completed.add((norm(row.get("aba")), int(row.get("linha") or 0)))
    return completed


def audit_headers(ws):
    headers = [cell.value for cell in ws[1]]
    return {header: index + 1 for index, header in enumerate(headers)}


def resolve_source_sheet(workbook, audit_sheet_name):
    wanted = norm(audit_sheet_name)
    for sheet in workbook.worksheets:
        if norm(sheet.title) == wanted:
            return sheet
    for sheet in workbook.worksheets:
        if wanted and (wanted in norm(sheet.title) or norm(sheet.title) in wanted):
            return sheet
    raise RuntimeError(f"aba da planilha original nao encontrada: {audit_sheet_name}")


def visible_text(driver):
    return norm(driver.execute_script("return document.body?.innerText || ''"))


def class_selection_visible(driver):
    text = visible_text(driver)
    if "voce esta em qual turma" in text or "selecione a turma" in text or "selecionar turma" in text:
        return True
    return False


def select_class_if_needed(driver, class_code):
    if not class_selection_visible(driver):
        return False
    g.select_created_account_class_safe(driver, class_code)
    return True


def audit_current(driver, class_code):
    try:
        return audit.audit_current_account(driver, class_code)
    except Exception as exc:
        return "ERRO", str(exc).splitlines()[0] or type(exc).__name__


def material_and_class_already_ok(driver, class_code, seconds=10):
    grade, letter, _ = g.class_parts(class_code)
    letter = letter.lower()
    deadline = time.time() + seconds
    last_text = ""
    while time.time() < deadline:
        text = visible_text(driver)
        header_titles = driver.execute_script(
            """
            return [...document.querySelectorAll('header [title], nav [title], button [title], [data-testid="tab-selector-Perfil"] [title], [aria-label="Perfil"] [title]')]
              .map((el) => el.getAttribute('title') || el.textContent || '')
              .join(' ');
            """
        )
        title_text = norm(header_titles)
        last_text = text
        compact = f"{text} {title_text}".replace(" ", "")
        compact_no_year = compact.replace("2026", "")
        class_tokens = {
            f"{grade}{letter}",
            f"{grade}o{letter}",
            f"{grade}º{letter}",
            f"{grade}ano{letter}",
            f"{grade}anos{letter}",
        }
        material_ok = (
            "material padrao 2026" in text
            or "material de 2026 ja esta aqui" in text
            or "visualizando o material padrao 2026" in text
            or "seu material de 2026 ja esta aqui" in text
        )
        class_ok = (
            f"{grade}{letter}2026" in compact
            or f"{grade}{letter}|2026" in compact
            or f"{grade}{letter}" in title_text.replace(" ", "")
            or (f"{grade}em" in compact and f"{grade}{letter}" in compact)
            or any(token in compact_no_year for token in class_tokens)
        )
        if material_ok and class_ok:
            print(f"Material 2026 e turma {grade}{letter.upper()} ja confirmados na tela.")
            return True
        time.sleep(0.5)
    print(f"Aviso: tela atual ainda nao confirmou material/turma. Amostra: {last_text[:160]}")
    return False


def profile_update_visible(driver):
    text = visible_text(driver)
    return "atualizacao das informacoes" in text


def finish_profile_update(driver, email, used_usernames):
    username = None
    for attempt in range(3):
        if not profile_update_visible(driver):
            return username
        print(f"Tela de atualizacao aberta; completando perfil antes de ativar material (tentativa {attempt + 1}).")
        updated = g.complete_profile_after_login(driver, str(email), used_usernames or [])
        if updated:
            username = updated
        time.sleep(2)
    if profile_update_visible(driver):
        raise RuntimeError("perfil ainda esta na tela de atualizacao; nao vou criar conta nova")
    return username


def repair_material_and_class(driver, code, class_code, email=None, used_usernames=None, retried_update=False):
    if profile_update_visible(driver):
        if not email:
            raise RuntimeError("tela de atualizacao aberta antes de ativar material")
        finish_profile_update(driver, email, used_usernames)

    if material_and_class_already_ok(driver, class_code):
        return

    if profile_update_visible(driver):
        if not email:
            raise RuntimeError("tela de atualizacao abriu antes de ativar material")
        finish_profile_update(driver, email, used_usernames)
        if material_and_class_already_ok(driver, class_code, seconds=3):
            return

    if select_class_if_needed(driver, class_code):
        return

    if not code:
        raise RuntimeError("sem codigo do material")
    try:
        if profile_update_visible(driver):
            if not email:
                raise RuntimeError("tela de atualizacao aberta antes de ativar material")
            finish_profile_update(driver, email, used_usernames)
        print("Ativando material pelo codigo da planilha para garantir a turma correta...")
        g.activate(driver, str(code), str(class_code))
        return
    except Exception as exc:
        exc_norm = norm(exc)
        if ("atualizacao das informacoes" in exc_norm or profile_update_visible(driver)) and not retried_update:
            if not email:
                raise
            print("Ativacao encontrou tela de atualizacao; completando perfil e tentando continuar.")
            finish_profile_update(driver, email, used_usernames)
            return repair_material_and_class(driver, code, class_code, email, used_usernames, retried_update=True)
        result, detail = audit_current(driver, class_code)
        if "material 2026 ja esta aberto" in exc_norm and ("turma encontrada" in norm(detail) or material_and_class_already_ok(driver, class_code, seconds=3)):
            print(f"Material ja estava aberto e auditoria confirmou: {detail}")
            return
        if "pede selecao" in norm(detail) and select_class_if_needed(driver, class_code):
            return
        raise RuntimeError(f"ativar material falhou: {str(exc).splitlines()[0] or type(exc).__name__}; auditoria: {detail}")


def already_on_one_home(driver):
    try:
        text = visible_text(driver)
        return "ativar novo material" in text or "minha semana" in text or "minhas materias" in text
    except Exception:
        return False


def reset_login_and_repair(driver, source_ws, columns, row, used_usernames):
    name = source_ws[f"{columns['name']}{row}"].value
    email = source_ws[f"{columns['email']}{row}"].value
    code = source_ws[f"{columns['code']}{row}"].value
    class_code = source_ws[f"{columns['class_code']}{row}"].value
    username = source_ws[f"{columns['username']}{row}"].value

    if not name or not email:
        raise RuntimeError("linha sem nome/email")
    if not code:
        raise RuntimeError("sem codigo do material")

    try:
        username, reset_link = g.search_lab(driver, str(name), str(email))
        g.reset_password(driver, reset_link, str(name))
        updated_username = g.login_one(driver, str(email), used_usernames)
        if updated_username:
            username = updated_username
        repair_material_and_class(driver, code, class_code, email, used_usernames)
    except g.AccountNotFound:
        try:
            print("Conta nao encontrada no Lab; tentando entrar no Geekie One antes de criar nova conta...")
            updated_username = g.login_one(driver, str(email), used_usernames)
            if updated_username:
                username = updated_username
            elif not username:
                username = g.username_from_email(email, used_usernames)
            repair_material_and_class(driver, code, class_code, email, used_usernames)
        except Exception as login_exc:
            if profile_update_visible(driver):
                print("Login abriu atualizacao de informacoes; completando perfil da conta existente.")
                updated_username = g.complete_profile_after_login(driver, str(email), used_usernames)
                if updated_username:
                    username = updated_username
                elif not username:
                    username = g.username_from_email(email, used_usernames)
                repair_material_and_class(driver, code, class_code, email, used_usernames)
                return username
            print(f"Login direto no Geekie One falhou; criando conta nova: {str(login_exc).splitlines()[0] or type(login_exc).__name__}")
            try:
                g.ensure_one_logged_out(driver)
            except Exception:
                pass
            try:
                username = g.create_one_account(
                    driver,
                    str(name),
                    str(email),
                    str(code),
                    str(class_code),
                    used_usernames,
                )
            except g.EmailAlreadyExists:
                print("E-mail ja existe no Geekie One; entrando na conta existente para finalizar.")
                try:
                    g.ensure_one_logged_out(driver)
                except Exception:
                    pass
                updated_username = g.login_one(driver, str(email), used_usernames)
                if updated_username:
                    username = updated_username
                elif not username:
                    username = g.username_from_email(email, used_usernames)
                repair_material_and_class(driver, code, class_code, email, used_usernames)

    return username


def login_and_fix_class_or_fallback(driver, source_ws, columns, row, used_usernames):
    email = source_ws[f"{columns['email']}{row}"].value
    code = source_ws[f"{columns['code']}{row}"].value
    class_code = source_ws[f"{columns['class_code']}{row}"].value

    try:
        updated_username = g.login_one(driver, str(email), used_usernames)
        repair_material_and_class(driver, code, class_code, email, used_usernames)
        return updated_username
    except Exception as exc:
        print(f"Aviso: correcao direta falhou; tentando fluxo completo com Lab: {str(exc).splitlines()[0] or type(exc).__name__}")
        try:
            g.logout(driver)
        except Exception:
            pass
        return reset_login_and_repair(driver, source_ws, columns, row, used_usernames)


def should_process(result, detail):
    detail_norm = norm(detail)
    if norm(result) == "ok":
        return False
    if "sem codigo" in detail_norm:
        return False
    return True


def selected_sheet_names(workbook, raw_indexes="", raw_names=""):
    names = set()
    if raw_indexes:
        for part in str(raw_indexes).split(","):
            part = part.strip()
            if part:
                names.add(workbook.worksheets[int(part) - 1].title)
    if raw_names:
        wanted = {norm(part) for part in str(raw_names).split(",") if part.strip()}
        for ws in workbook.worksheets:
            if norm(ws.title) in wanted:
                names.add(ws.title)
    return {norm(name) for name in names}


def prefer_direct_class_fix(detail):
    detail_norm = norm(detail)
    return (
        "pede selecao de turma" in detail_norm
        or "turma esperada nao confirmada" in detail_norm
        or "turma encontrada" in detail_norm
        or "stale element" in detail_norm
    )


def run(args):
    workbook_path = Path(args.workbook).resolve()
    audit_path = Path(args.audit).resolve()
    source_wb = load_workbook(workbook_path)
    audit_wb = load_workbook(audit_path, data_only=True)
    audit_ws = audit_wb.active
    audit_idx = audit_headers(audit_ws)
    completed = set() if args.redo else completed_from_log()
    selected_sheets = selected_sheet_names(source_wb, args.only_sheet_indexes, args.only_sheets)

    tasks = []
    for audit_row in range(2, audit_ws.max_row + 1):
        result = audit_ws.cell(audit_row, audit_idx["resultado"]).value
        detail = audit_ws.cell(audit_row, audit_idx["detalhe"]).value
        if not should_process(result, detail):
            continue
        source_sheet_name = audit_ws.cell(audit_row, audit_idx["aba"]).value
        if selected_sheets and norm(source_sheet_name) not in selected_sheets:
            continue
        source_row = audit_ws.cell(audit_row, audit_idx["linha"]).value
        if (norm(source_sheet_name), int(source_row or 0)) in completed:
            continue
        tasks.append((audit_row, source_sheet_name, int(source_row), result, detail))

    if args.limit:
        tasks = tasks[: args.limit]

    print(f"Pendencias corrigiveis encontradas: {len(tasks)}")
    if args.dry_run:
        for _, sheet_name, row, result, detail in tasks:
            print(f"[DRY] {sheet_name} linha {row}: {result} - {detail}")
        return

    sheet_cache = {}
    used_usernames = []
    for ws in source_wb.worksheets:
        try:
            first_data_row, columns = g.detect_layout(ws)
        except Exception:
            continue
        username_col = columns.get("username")
        if username_col:
            for row in range(first_data_row, ws.max_row + 1):
                value = ws[f"{username_col}{row}"].value
                if value:
                    used_usernames.append(value)

    driver = g.new_driver()
    try:
        g.ensure_logins(driver)
        for index, (audit_row, sheet_name, row, _result, detail) in enumerate(tasks, start=1):
            source_ws = sheet_cache.get(sheet_name)
            if source_ws is None:
                source_ws = resolve_source_sheet(source_wb, sheet_name)
                first_data_row, columns = g.detect_layout(source_ws)
                sheet_cache[sheet_name] = source_ws
            else:
                first_data_row, columns = g.detect_layout(source_ws)

            name = source_ws[f"{columns['name']}{row}"].value
            email = source_ws[f"{columns['email']}{row}"].value
            code = source_ws[f"{columns['code']}{row}"].value
            username_cell = f"{columns['username']}{row}"
            print(f"[{index}/{len(tasks)}] {source_ws.title} linha {row}: {name}")

            if not name or not email:
                append_log([source_ws.title, row, name, "PULADO", "linha sem nome/email"])
                continue
            if not code:
                append_log([source_ws.title, row, name, "PULADO", "sem codigo do material"])
                continue

            try:
                if prefer_direct_class_fix(detail):
                    username = login_and_fix_class_or_fallback(driver, source_ws, columns, row, used_usernames)
                else:
                    username = reset_login_and_repair(driver, source_ws, columns, row, used_usernames)

                if username:
                    source_ws[username_cell] = username
                    if username not in used_usernames:
                        used_usernames.append(username)
                if columns.get("status"):
                    source_ws[f"{columns['status']}{row}"] = "ok"
                audit_ws.cell(audit_row, audit_idx["resultado"]).value = "OK"
                audit_ws.cell(audit_row, audit_idx["detalhe"]).value = "corrigido pelo script"
                source_wb.save(workbook_path)
                audit_wb.save(audit_path)
                append_log([source_ws.title, row, name, "OK", username or source_ws[username_cell].value or ""])
                print(f"OK: usuario={source_ws[username_cell].value or ''}; status=ok salvo na planilha original")
            except Exception as exc:
                lines = [line.strip() for line in str(exc).splitlines() if line.strip()]
                detail_text = lines[0] if lines else type(exc).__name__
                if detail_text.lower() in {"message:", "message"}:
                    detail_text = type(exc).__name__
                append_log([source_ws.title, row, name, "ERRO", detail_text])
                print(f"ERRO: {detail_text}")
                try:
                    driver.save_screenshot(str(Path(__file__).with_name(f"corrigir_erro_linha_{row}.png")))
                    Path(__file__).with_name(f"corrigir_erro_linha_{row}.html").write_text(driver.page_source, encoding="utf-8")
                except Exception:
                    pass
                if args.stop_on_error:
                    raise
            finally:
                try:
                    g.logout(driver)
                except Exception:
                    try:
                        g.ensure_one_logged_out(driver)
                    except Exception:
                        try:
                            driver.get(g.ONE_URL)
                        except Exception:
                            pass
                time.sleep(0.5)
    finally:
        driver.quit()
        source_wb.save(workbook_path)


def main():
    parser = argparse.ArgumentParser(description="Corrige alunos pendentes apontados pela auditoria Geekie.")
    parser.add_argument("workbook")
    parser.add_argument("--audit", default="auditoria_geekie_resultado.xlsx")
    parser.add_argument("--limit", type=int, default=0, help="Processa apenas as primeiras N pendencias corrigiveis.")
    parser.add_argument("--dry-run", action="store_true")
    parser.add_argument("--stop-on-error", action="store_true")
    parser.add_argument("--redo", action="store_true", help="Ignora o log e tenta novamente linhas ja marcadas como OK.")
    parser.add_argument("--only-sheet-indexes", default="", help="Processa apenas abas da planilha original por indice, ex: 4 para 9 ano.")
    parser.add_argument("--only-sheets", default="", help="Processa apenas abas por nome, separadas por virgula.")
    args = parser.parse_args()
    run(args)


if __name__ == "__main__":
    main()
