from __future__ import annotations

import argparse
import subprocess
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parent
RESULTS = ROOT / "resultados"


def choose(prompt: str, options: list[tuple[str, str]], default: str | None = None) -> str:
    print()
    print(prompt)
    for key, label in options:
        suffix = " (padrao)" if default and key == default else ""
        print(f"  {key} - {label}{suffix}")
    while True:
        value = input("Opcao: ").strip() or (default or "")
        if any(key == value for key, _ in options):
            return value
        print("Opcao invalida.")


def list_workbooks() -> list[Path]:
    files = sorted(
        [path for path in ROOT.glob("*.xlsx") if not path.name.lower().startswith("~$")],
        key=lambda path: (path.name.lower().startswith("auditoria_"), path.name.lower()),
    )
    return files


def choose_workbook() -> Path:
    files = list_workbooks()
    if not files:
        raise SystemExit("Nenhuma planilha .xlsx encontrada nesta pasta.")
    if len(files) == 1:
        print(f"Planilha encontrada: {files[0].name}")
        return files[0]
    options = [(str(index + 1), path.name) for index, path in enumerate(files)]
    selected = choose("Escolha a planilha:", options, "1")
    return files[int(selected) - 1]


def norm(value: object) -> str:
    text = "" if value is None else str(value)
    return " ".join(text.strip().lower().split())


def detect_geekie_sheets(workbook: Path) -> list[tuple[int, str]]:
    wb = load_workbook(workbook, read_only=True, data_only=True)
    candidates: list[tuple[int, str]] = []
    for index, ws in enumerate(wb.worksheets, start=1):
        headers = {norm(cell.value) for cell in next(ws.iter_rows(min_row=1, max_row=1), [])}
        title = norm(ws.title)
        if title.startswith("geral"):
            continue
        looks_like_geekie = "geekie" in title or {"nomealuno", "loginoffice365", "codigo do material"} & headers
        if looks_like_geekie:
            candidates.append((index, ws.title))
    wb.close()
    return candidates


def choose_sheets(workbook: Path, allow_many: bool) -> list[int]:
    candidates = detect_geekie_sheets(workbook)
    if not candidates:
        raise SystemExit("Nao encontrei abas Geekie nessa planilha.")

    options = [(str(i), f"{title} (aba {i})") for i, title in candidates]
    if allow_many:
        options.insert(0, ("T", "Todas as abas detectadas"))
    options.append(("M", "Digitar indice manualmente"))

    selected = choose("Escolha qual serie/aba fazer:", options, "T" if allow_many else options[0][0])
    if selected == "T":
        return [index for index, _ in candidates]
    if selected == "M":
        raw = input("Digite o(s) indice(s), separados por virgula: ").strip()
        return [int(part.strip()) for part in raw.split(",") if part.strip()]
    return [int(selected)]


def ask_start() -> int:
    raw = input("Comecar em qual linha/aluno? (padrao 2): ").strip()
    return int(raw or "2")


def latest_audit_file() -> Path | None:
    candidates = list(ROOT.glob("auditoria_geekie_resultado*.xlsx")) + list(RESULTS.glob("auditoria_geekie_resultado*.xlsx"))
    if not candidates:
        return None
    return max(candidates, key=lambda path: path.stat().st_mtime)


def run_command(args: list[str]) -> int:
    print()
    print("Executando:")
    print(" ".join(f'"{arg}"' if " " in arg else arg for arg in args))
    print()
    return subprocess.call(args, cwd=ROOT)


def process_run(workbook: Path) -> int:
    sheets = choose_sheets(workbook, allow_many=False)
    start = ask_start()
    return run_command([
        sys.executable,
        "geekie_automation.py",
        "run",
        str(workbook),
        "--sheet-index",
        str(sheets[0]),
        "--start",
        str(start),
        "--continue-on-error",
    ])


def process_audit(workbook: Path) -> int:
    RESULTS.mkdir(exist_ok=True)
    sheets = choose_sheets(workbook, allow_many=True)
    start = ask_start()
    stamp = datetime.now().strftime("%Y%m%d_%H%M")
    output = RESULTS / f"auditoria_geekie_resultado_{stamp}.xlsx"
    return run_command([
        sys.executable,
        "auditoria_geekie.py",
        str(workbook),
        "--sheet-indexes",
        ",".join(map(str, sheets)),
        "--start",
        str(start),
        "--output",
        str(output),
    ])


def process_correct(workbook: Path) -> int:
    sheets = choose_sheets(workbook, allow_many=True)
    audit = latest_audit_file()
    if audit:
        use_latest = choose("Auditoria para usar:", [("1", audit.name), ("2", "Digitar caminho manualmente")], "1")
        if use_latest == "2":
            audit = Path(input("Caminho da auditoria: ").strip().strip('"'))
    else:
        audit = Path(input("Caminho da auditoria: ").strip().strip('"'))
    return run_command([
        sys.executable,
        "corrigir_pendentes.py",
        str(workbook),
        "--audit",
        str(audit),
        "--only-sheet-indexes",
        ",".join(map(str, sheets)),
    ])


def main() -> None:
    parser = argparse.ArgumentParser(description="Menu Geekie")
    parser.add_argument("mode", choices=["run", "audit", "correct", "list"])
    args = parser.parse_args()

    workbook = choose_workbook()
    if args.mode == "list":
        print(f"Planilha: {workbook.name}")
        for index, title in detect_geekie_sheets(workbook):
            print(f"{index}: {title}")
        code = 0
    elif args.mode == "run":
        code = process_run(workbook)
    elif args.mode == "audit":
        code = process_audit(workbook)
    else:
        code = process_correct(workbook)
    raise SystemExit(code)


if __name__ == "__main__":
    main()
